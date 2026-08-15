"""运行时生成 Excel 导入模板（无需随包分发静态文件）。

openpyxl 为函数内延迟导入，避免拖慢应用启动。
"""

_HEADER_FILL_ARGS = ("solid", "DDEBF7")

_FILAMENT_HEADERS = ["编号", "批次", "断裂强力", "强力单位", "直径(μm)", "密度(g/cm³)", "备注", "测试日期"]
_TEXTILE_HEADERS = ["编号", "批次", "断裂强力", "强力单位", "线密度(dtex)", "密度(g/cm³)", "备注", "测试日期"]

_FILAMENT_EXAMPLES = [
    ["F-01", "碳纤维-A", 10.0, "cN", 7.0, None, "示例数据，请删除后填写", "2026-08-14"],
    ["F-02", "碳纤维-A", 9.5, "cN", 7.2, None, None, None],
    ["F-03", "碳纤维-A", 10.8, "cN", 6.8, None, None, None],
]
_TEXTILE_EXAMPLES = [
    ["T-01", "涤纶-B", 12.5, "cN", 1.5, 1.38, "示例数据，请删除后填写", "2026-08-14"],
    ["T-02", "涤纶-B", 11.8, "cN", 1.55, 1.38, None, None],
    ["T-03", "涤纶-B", 13.1, "cN", 1.48, 1.38, None, None],
]

_INSTRUCTIONS = [
    ("导入模板填写说明", True),
    ("", False),
    ("1. 本文件包含「单丝模式示例」与「纺织模式示例」两个示例表，请按需选择其中一个，", False),
    ("   另存为新文件后填写数据再导入。", False),
    ("2. 断裂强力为必填列；强力单位可留空（默认 cN），支持 cN / N / mN。", False),
    ("3. 单丝模式（碳纤维/玻璃纤维等单丝，按 σ = 4F/(πd²) 计算，结果 MPa/GPa）：", False),
    ("   直径(μm) 为必填列；密度列可删除。", False),
    ("4. 纺织模式（化纤/棉等，断裂强度 = F/线密度，结果 cN/dtex）：", False),
    ("   线密度(dtex) 为必填列；密度(g/cm³) 可选，填写后报告可附 MPa 换算值。", False),
    ("5. 测试日期格式：2026-08-14 或 2026/8/14，可留空。", False),
    ("6. 文件可另存为 .xlsx / .csv / .txt 后导入（CSV 编码支持 UTF-8 与 GBK）。", False),
    ("7. 编号、批次、备注等文本列可留空；整行留空会自动跳过。", False),
]

_COL_WIDTHS = {"编号": 10, "批次": 12, "断裂强力": 12, "强力单位": 10,
               "直径(μm)": 10, "线密度(dtex)": 13, "密度(g/cm³)": 12, "备注": 24, "测试日期": 12}


def _fill_sheet(ws, headers, examples, styles) -> None:
    from openpyxl.utils import get_column_letter

    header_fill, header_font, alignment = styles
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = alignment(horizontal="center")
    for row in examples:
        ws.append(row)
    for i, name in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = _COL_WIDTHS.get(name, 12)
    ws.freeze_panes = "A2"


def generate_template(path) -> None:
    """生成模板 xlsx：两个模式示例 sheet + 填写说明 sheet。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    styles = (PatternFill("solid", fgColor="DDEBF7"), Font(bold=True), Alignment)
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "单丝模式示例"
    _fill_sheet(ws1, _FILAMENT_HEADERS, _FILAMENT_EXAMPLES, styles)

    ws2 = wb.create_sheet("纺织模式示例")
    _fill_sheet(ws2, _TEXTILE_HEADERS, _TEXTILE_EXAMPLES, styles)

    ws3 = wb.create_sheet("填写说明")
    for text, bold in _INSTRUCTIONS:
        ws3.append([text])
        if bold:
            ws3.cell(row=ws3.max_row, column=1).font = Font(bold=True, size=13)
    ws3.column_dimensions["A"].width = 88
    wb.save(path)
