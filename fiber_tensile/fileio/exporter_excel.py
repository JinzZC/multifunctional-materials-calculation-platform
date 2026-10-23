"""Excel 报告导出：统计汇总 / 数据明细 / Weibull 散点 / Weibull 拟合直线。

openpyxl/matplotlib 均为函数内延迟导入，避免拖慢应用启动。
"""

from ..core.models import MODE_LABELS
from ..core.modes import get_mode_spec
from ..core.weibull import (
    rank_method_label, weibull_in_gpa, weibull_intercept, weibull_line_data, weibull_x_label,
)
from .excel_safe import set_cell_value
from .atomic import atomic_write


def _styles():
    """延迟构建 openpyxl 样式对象。"""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    return (
        PatternFill("solid", fgColor="DDEBF7"),
        Font(bold=True),
        Border(left=Side(style="thin", color="B0B0B0"),
               right=Side(style="thin", color="B0B0B0"),
               top=Side(style="thin", color="B0B0B0"),
               bottom=Side(style="thin", color="B0B0B0")),
        Alignment,
    )


def _write_kv(ws, row: int, key, value, styles) -> int:
    _, bold, _, _ = styles
    ws.cell(row, 1, key).font = bold
    set_cell_value(ws.cell(row, 2), value)
    return row + 1


def _write_table(ws, row: int, headers: list, rows: list, styles) -> int:
    hdr_fill, bold, border, alignment = styles
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row, col, h)
        c.font = bold
        c.fill = hdr_fill
        c.border = border
        c.alignment = alignment(horizontal="center")
    row += 1
    for r in rows:
        for col, v in enumerate(r, start=1):
            c = set_cell_value(ws.cell(row, col), v)
            c.border = border
            if isinstance(v, (int, float)):
                c.number_format = "0.00"
        row += 1
    return row


def export_excel_report(path, batch, analysis=None, weibull=None) -> None:
    """导出一个批次的 Excel 报告。analysis/weibull 由调用方计算后传入。"""
    from io import BytesIO

    import matplotlib.pyplot as plt
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    from ..plotting import build_weibull_figure

    styles = _styles()
    spec = get_mode_spec(batch.mode)
    unit = spec.strength_unit
    weibull_error = "无法进行 Weibull 分析（至少需要 2 个有效且不完全相同的强度值）"
    if weibull is not None:
        try:
            weibull = weibull_in_gpa(weibull, unit, batch.density)
        except ValueError as exc:
            weibull_error = str(exc)
            weibull = None
    wb = Workbook()
    ws = wb.active
    ws.title = "统计汇总"
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 14

    row = 1
    c = ws.cell(row, 1, f"纤维拉伸强度分析报告 — {batch.batch_name}")
    c.font = Font(bold=True, size=14)
    row += 2
    c = ws.cell(row, 1, "批次信息")
    c.font = Font(bold=True, size=12)
    row += 1
    info = [
        ("批次名称", batch.batch_name),
        ("计算模式", MODE_LABELS[batch.mode]),
        ("纤维类型", batch.fiber_type or "—"),
        ("密度(g/cm³)", batch.density if batch.density is not None else "—"),
        ("备注", batch.remark or "—"),
        ("创建时间", batch.created_at or "—"),
        ("样本数", batch.sample_count),
    ]
    for k, v in info:
        row = _write_kv(ws, row, k, v, styles)
    row += 1

    c = ws.cell(row, 1, "统计结果")
    c.font = Font(bold=True, size=12)
    row += 1
    if analysis is not None:
        headers = ["样本数 N", f"均值({unit})", f"标准差({unit})", "变异系数 CV%",
                   f"最小值({unit})", f"最大值({unit})"]
        values = [
            analysis.n,
            round(analysis.mean, 2),
            round(analysis.std, 2) if analysis.std is not None else "—",
            round(analysis.cv_percent, 2) if analysis.cv_percent is not None else "—",
            round(analysis.min_val, 2),
            round(analysis.max_val, 2),
        ]
        row = _write_table(ws, row, headers, [values], styles)
    else:
        ws.cell(row, 1, "无数据")
        row += 1
    row += 1

    c = ws.cell(row, 1, "Weibull 分析结果")
    c.font = Font(bold=True, size=12)
    row += 1
    if weibull is not None:
        kv_rows = [
            ("排序概率方法", rank_method_label(weibull.rank_method)),
            ("Weibull 模数 m", round(weibull.m, 4)),
            ("特征强度 σ₀(GPa)", weibull.sigma0),
            ("位置参数 γ(GPa)", weibull.gamma),
            ("横坐标 X", weibull_x_label(weibull)),
            ("拟合直线斜率 m", weibull.m),
            ("拟合直线截距 a", weibull_intercept(weibull)),
            ("拟合优度 R²", round(weibull.r2, 4)),
        ]
        for k, v in kv_rows:
            row = _write_kv(ws, row, k, v, styles)
        fig = build_weibull_figure([weibull], labels=[batch.batch_name], sigma0_unit="GPa")
        buf = BytesIO()
        fig.savefig(buf, format="png", dpi=150)
        plt.close(fig)
        img = XLImage(buf)
        img.width = 640
        img.height = 470
        ws.add_image(img, "E2")
    else:
        ws.cell(row, 1, weibull_error)

    # ---- Sheet2 数据明细 ----
    ws2 = wb.create_sheet("数据明细")
    dim_col = spec.dim_label
    headers2 = ["编号", "断裂强力", "强力单位", dim_col, "密度(g/cm³)",
                f"强度({unit})", "强度(MPa)", "备注", "测试日期"]
    rows2 = []
    for s in batch.samples:
        dim = getattr(s, spec.dim_field)
        rows2.append([
            s.sample_no or "—",
            s.breaking_force,
            s.force_unit,
            dim,
            batch.density if batch.density is not None else "—",
            round(s.strength, 2) if s.strength is not None else "—",
            round(s.strength_mpa, 2) if s.strength_mpa is not None else "—",
            s.remark or "",
            s.test_date or "",
        ])
    _write_table(ws2, 1, headers2, rows2, styles)
    ws2.freeze_panes = "A2"
    for i, w in enumerate([10, 11, 9, 13, 12, 13, 12, 22, 12], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # ---- Sheet3 Weibull 分析数据 ----
    if weibull is not None:
        ws3 = wb.create_sheet("Weibull 分析数据")
        x_label = weibull_x_label(weibull)
        intercept = weibull_intercept(weibull)
        headers3 = ["序号 i", "σ(GPa)", f"F̂({rank_method_label(weibull.rank_method)})",
                    x_label, "ln(−ln(1−F̂))", "拟合 Y = m·X+a"]
        rows3 = []
        for i, (s_val, f_hat, x, y) in enumerate(zip(
                weibull.strengths_sorted, weibull.median_ranks,
                weibull.x, weibull.y), start=1):
            rows3.append([i, s_val, f_hat, x, y, weibull.m * x + intercept])
        _write_table(ws3, 1, headers3, rows3, styles)
        ws3.freeze_panes = "A2"
        for i, w in enumerate([8, 20, 44, 24, 24, 24], start=1):
            ws3.column_dimensions[get_column_letter(i)].width = w
        for cells in ws3.iter_rows(min_row=2, min_col=2):
            for cell in cells:
                cell.number_format = "0.0000000000"

        # 独立的 X/Y 数值列，直接在 Origin 中指定 A(X)、B(Y) 后画线。
        ws4 = wb.create_sheet("Weibull 拟合直线数据")
        line_x, line_y = weibull_line_data(weibull)
        _write_table(ws4, 1, [f"拟合 X = {x_label}", "拟合 Y = m·X+a"],
                     list(zip(line_x, line_y)), styles)
        ws4.freeze_panes = "A2"
        for column in ("A", "B"):
            ws4.column_dimensions[column].width = 28
        for cells in ws4.iter_rows(min_row=2, max_col=2):
            for cell in cells:
                cell.number_format = "0.0000000000"
        parameters = [
            ("横坐标 X", x_label),
            ("纵坐标 Y", "ln(−ln(1−F))"),
            ("斜率 m", weibull.m),
            ("截距 a", intercept),
            ("尺度参数 σ₀(GPa)", weibull.sigma0),
            ("位置参数 γ(GPa)", weibull.gamma),
            ("拟合优度 R²", weibull.r2),
            ("排序概率方法", rank_method_label(weibull.rank_method)),
            ("拟合直线点数", len(line_x)),
            ("绘图范围", "散点 X 最小值至最大值，X 等间距"),
            ("Origin 绘图", "本表 A/B 列画拟合线；分析数据表 D/E 列画散点，均用线性坐标轴。"),
        ]
        for r, (key, value) in enumerate(parameters, start=1):
            ws4.cell(r, 4, key).font = styles[1]
            cell = set_cell_value(ws4.cell(r, 5), value)
            if isinstance(value, float):
                cell.number_format = "0.0000000000"
        ws4.column_dimensions["D"].width = 24
        ws4.column_dimensions["E"].width = 82

    atomic_write(path, wb.save)
