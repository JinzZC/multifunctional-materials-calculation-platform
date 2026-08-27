"""统一批量计算服务（无 Qt 依赖）。

每个文件独立解析、计算和保存；某个文件失败不会回滚其他已成功文件。
调用方可通过 ``cancelled`` 和 ``progress`` 实现后台取消与进度显示。
"""

from __future__ import annotations

import csv
import math
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Iterable

from ..core.models import Batch
from ..curve_analysis.core import analyze_curve
from ..curve_analysis.importer import parse_curve_file
from ..curve_analysis.repository import create_curve_analysis
from ..data.repository import create_batch_with_samples
from ..fileio.importer import parse_file
from ..surface_energy.core import METHODS
from ..surface_energy.importer import parse_measurement_file
from ..surface_energy.repository import create_measurement

SUPPORTED_DATA_FILES = {".xlsx", ".xlsm", ".csv", ".txt", ".tsv"}
DATA_TYPES = {"fiber", "surface", "curve"}


@dataclass(frozen=True)
class BatchOptions:
    data_type: str
    recursive: bool = True
    duplicate_policy: str = "rename"  # rename / skip / allow
    fiber_mode: str | None = None
    surface_method: str = "owrk"
    curve_diameter_um: float = 5.4
    curve_gauge_mm: float = 25.0
    curve_fit_lo: float = 0.2
    curve_fit_hi: float = 0.7
    curve_disp_unit: str = "mm"
    curve_force_unit: str = "N"

    def __post_init__(self) -> None:
        if self.data_type not in DATA_TYPES:
            raise ValueError(f"未知批量数据类型：{self.data_type!r}")
        if self.duplicate_policy not in {"rename", "skip", "allow"}:
            raise ValueError("重复记录策略必须是 rename/skip/allow")
        if self.surface_method not in METHODS:
            raise ValueError(f"未知表面能方法：{self.surface_method!r}")
        if self.data_type == "curve":
            if not math.isfinite(self.curve_diameter_um) or self.curve_diameter_um <= 0:
                raise ValueError("曲线批处理的纤维直径必须为正数")
            if not math.isfinite(self.curve_gauge_mm) or self.curve_gauge_mm <= 0:
                raise ValueError("曲线批处理的标距必须为正数")
            if not 0 < self.curve_fit_lo < self.curve_fit_hi < 1:
                raise ValueError("曲线拟合区间必须满足 0 < 下限 < 上限 < 1")
            if self.curve_disp_unit not in {"mm", "um"}:
                raise ValueError("位移单位仅支持 mm/μm")
            if self.curve_force_unit not in {"N", "cN", "mN"}:
                raise ValueError("力单位仅支持 N/cN/mN")


@dataclass
class BatchResult:
    source: str
    data_type: str
    status: str
    record_name: str = ""
    record_id: int | None = None
    input_rows: int = 0
    valid_rows: int = 0
    warning_count: int = 0
    summary: str = ""
    error: str = ""
    warnings: list[str] = field(default_factory=list)


def discover_data_files(paths: Iterable[str | Path], recursive: bool = True) -> list[Path]:
    """从文件和文件夹混合输入中发现可处理文件，去重后自然排序。"""
    found: dict[str, Path] = {}
    for raw in paths:
        path = Path(raw)
        candidates = [path] if path.is_file() else (
            path.rglob("*") if recursive and path.is_dir()
            else path.glob("*") if path.is_dir() else ())
        for candidate in candidates:
            if candidate.is_file() and candidate.suffix.lower() in SUPPORTED_DATA_FILES:
                resolved = candidate.resolve()
                found[str(resolved).casefold()] = resolved
    return sorted(found.values(), key=lambda p: (p.name.casefold(), str(p).casefold()))


def _existing_name(conn, data_type: str, name: str) -> bool:
    table, column = {
        "fiber": ("batches", "batch_name"),
        "surface": ("se_measurements", "sample_name"),
        "curve": ("fiber_curves", "sample_name"),
    }[data_type]
    return conn.execute(
        f"SELECT 1 FROM {table} WHERE {column} = ? LIMIT 1", (name,)
    ).fetchone() is not None


def _resolved_name(conn, data_type: str, requested: str, policy: str) -> str | None:
    name = requested.strip() or "未命名记录"
    if not _existing_name(conn, data_type, name) or policy == "allow":
        return name
    if policy == "skip":
        return None
    index = 2
    while _existing_name(conn, data_type, f"{name} ({index})"):
        index += 1
    return f"{name} ({index})"


def _warning_text(item) -> str:
    row = getattr(item, "row_number", None)
    prefix = f"第 {row} 行：" if row else ""
    return prefix + str(getattr(item, "message", item))


def _process_fiber(conn, path: Path, options: BatchOptions) -> BatchResult:
    preview = parse_file(path, mode=options.fiber_mode)
    if not preview.samples:
        raise ValueError("没有可导入的有效纤维数据")
    requested = preview.batch_from_file or path.stem
    name = _resolved_name(conn, "fiber", requested, options.duplicate_policy)
    if name is None:
        return BatchResult(str(path), "fiber", "跳过", requested,
                           input_rows=preview.total_rows,
                           valid_rows=preview.valid_rows,
                           warning_count=len(preview.warnings),
                           summary="存在同名批次，已按设置跳过")
    densities = sorted({s.density_gcm3 for s in preview.samples
                        if s.density_gcm3 is not None})
    density = densities[0] if densities else None
    warnings = [_warning_text(w) for w in preview.warnings]
    if len(densities) > 1:
        warnings.append(f"文件含多种密度，批次密度采用 {density:g} g/cm³")
    bid, count = create_batch_with_samples(
        conn, Batch(batch_name=name, mode=preview.inferred_mode, density=density,
                    remark=f"批量计算中心导入：{path.name}"), preview.samples)
    return BatchResult(
        str(path), "fiber", "成功", name, bid, preview.total_rows, count,
        len(warnings), f"已创建纤维批次，写入 {count} 条数据", warnings=warnings)


def _process_surface(conn, path: Path, options: BatchOptions) -> BatchResult:
    preview = parse_measurement_file(path)
    if preview.errors:
        details = "；".join(_warning_text(e) for e in preview.errors[:5])
        raise ValueError(f"存在 {len(preview.errors)} 条阻断性错误：{details}")
    method = preview.method or options.surface_method
    requested = preview.sample_name or path.stem
    name = _resolved_name(conn, "surface", requested, options.duplicate_policy)
    if name is None:
        return BatchResult(str(path), "surface", "跳过", requested,
                           input_rows=len(preview.points), valid_rows=len(preview.points),
                           warning_count=len(preview.warnings),
                           summary="存在同名表面能记录，已按设置跳过")
    mid = create_measurement(conn, name, method, preview.points,
                             preview.remark, preview.test_date)
    warnings = [_warning_text(w) for w in preview.warnings]
    return BatchResult(
        str(path), "surface", "成功", name, mid, len(preview.points),
        len(preview.points), len(warnings),
        f"已按 {METHODS[method].label} 计算并保存", warnings=warnings)


def _process_curve(conn, path: Path, options: BatchOptions) -> BatchResult:
    disp, force = parse_curve_file(path)
    disp_scale = 0.001 if options.curve_disp_unit == "um" else 1.0
    force_scale = {"N": 1.0, "cN": 0.01, "mN": 0.001}[options.curve_force_unit]
    disp_mm = [value * disp_scale for value in disp]
    force_n = [value * force_scale for value in force]
    result = analyze_curve(
        disp_mm, force_n, options.curve_diameter_um, options.curve_gauge_mm,
        options.curve_fit_lo, options.curve_fit_hi)
    requested = path.stem
    name = _resolved_name(conn, "curve", requested, options.duplicate_policy)
    if name is None:
        return BatchResult(str(path), "curve", "跳过", requested,
                           input_rows=len(disp), valid_rows=len(disp),
                           summary="存在同名曲线记录，已按设置跳过")
    cid = create_curve_analysis(
        conn, name, options.curve_diameter_um, options.curve_gauge_mm, result,
        disp_mm, force_n, options.curve_fit_lo, options.curve_fit_hi,
        remark=f"批量计算中心导入：{path.name}")
    return BatchResult(
        str(path), "curve", "成功", name, cid, len(disp), len(disp), 0,
        f"E={result.e_gpa:.3f} GPa，σ={result.strength_mpa:.2f} MPa")


_PROCESSORS = {
    "fiber": _process_fiber,
    "surface": _process_surface,
    "curve": _process_curve,
}


def process_files(conn, paths: Iterable[str | Path], options: BatchOptions,
                  cancelled: Callable[[], bool] | None = None,
                  progress: Callable[[int, int, BatchResult], None] | None = None
                  ) -> list[BatchResult]:
    """逐文件计算并保存；失败转为结果行，不中断整批。"""
    files = [Path(path) for path in paths]
    results: list[BatchResult] = []
    for index, path in enumerate(files, 1):
        if cancelled and cancelled():
            break
        try:
            result = _PROCESSORS[options.data_type](conn, path, options)
        except Exception as exc:
            result = BatchResult(str(path), options.data_type, "失败",
                                 record_name=path.stem, error=str(exc),
                                 summary="计算或保存失败")
        results.append(result)
        if progress:
            progress(index, len(files), result)
    return results


def export_results_csv(path: str | Path, results: Iterable[BatchResult]) -> None:
    """导出带 UTF-8 BOM 的批处理审计日志。"""
    target = Path(path)
    with target.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["源文件", "数据类型", "状态", "记录名称", "记录ID",
                         "输入行数", "有效行数", "警告数", "摘要", "错误", "警告明细"])
        for result in results:
            writer.writerow([
                result.source, result.data_type, result.status, result.record_name,
                result.record_id or "", result.input_rows, result.valid_rows,
                result.warning_count, result.summary, result.error,
                "\n".join(result.warnings),
            ])


def export_results_excel(path: str | Path, results: Iterable[BatchResult]) -> None:
    """导出批处理结果与警告明细为 Excel 审计表。"""
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    from ..fileio.atomic import atomic_write_bytes
    from ..fileio.excel_safe import append_safe_row

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "批量计算结果"
    headers = ["源文件", "数据类型", "状态", "记录名称", "记录ID",
               "输入行数", "有效行数", "警告数", "摘要", "错误", "警告明细"]
    append_safe_row(sheet, headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
    for result in results:
        append_safe_row(sheet, [
            result.source, result.data_type, result.status, result.record_name,
            result.record_id or "", result.input_rows, result.valid_rows,
            result.warning_count, result.summary, result.error,
            "\n".join(result.warnings),
        ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = (44, 12, 10, 24, 10, 12, 12, 10, 40, 50, 60)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    stream = BytesIO()
    workbook.save(stream)
    atomic_write_bytes(path, stream.getvalue())
