"""Excel 批量提取与合并的纯文件处理层。

列按规范化后的表头名称对齐，首个出现的原始表头作为输出名称。每一行都可附带
来源文件夹、文件、工作表和原始行号，保证合并结果可以回溯到原始数据。
"""

from __future__ import annotations

import os
import csv
import re
import tempfile
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

SUPPORTED_SUFFIXES = {".xlsx", ".xlsm", ".csv", ".txt", ".tsv"}
DELIMITED_SUFFIXES = {".csv", ".txt", ".tsv"}
SOURCE_HEADERS = ("来源文件夹", "来源文件", "来源工作表", "原始行号")
EXCEL_MAX_ROWS = 1_048_576
SAMPLE_NAME_ALIASES = {
    "样品", "样品名", "样品名称", "试样", "试样名", "试样名称", "样本名", "样本名称",
    "sample", "samplename", "sample_name", "specimen", "specimenname",
    "specimen_name",
}


class MergeCancelled(Exception):
    """用户主动取消批量提取。"""


@dataclass(frozen=True)
class MergeOptions:
    header_row: int = 1
    sheet_name: str | None = None       # None 表示全部工作表
    sheet_match_mode: str = "exact"     # exact / contains；仅 Excel 有工作表
    recursive: bool = True
    include_source: bool = True
    column_mode: str = "union"          # union / intersection
    skip_empty_rows: bool = True
    sort_by_sample: bool = False
    sample_column: str | None = None     # None 表示按常见别名自动识别
    formula_mode: str = "values"         # values（优先缓存值）/ formulas（保留公式文本）
    column_mappings: tuple[tuple[str, str], ...] = ()
    sort_columns: tuple[str, ...] = ()
    duplicate_mode: str = "all"           # all / first / last
    duplicate_columns: tuple[str, ...] = ()


@dataclass
class FileStatus:
    path: str
    sheet: str
    rows: int = 0
    columns: int = 0
    status: str = "成功"
    message: str = ""


@dataclass
class MergeResult:
    headers: list[str] = field(default_factory=list)
    rows: list[list] = field(default_factory=list)
    statuses: list[FileStatus] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def successful_sheets(self) -> int:
        return sum(s.status == "成功" for s in self.statuses)

    @property
    def failed_sheets(self) -> int:
        return sum(s.status != "成功" for s in self.statuses)


def discover_excel_files(folders: Iterable[str | os.PathLike], recursive: bool = True,
                         excluded_paths: Iterable[str | os.PathLike] = (),
                         filename_contains: str | None = None) -> list[Path]:
    """递归扫描并去重表格文件，可按文件名包含文字筛选。"""
    excluded = {_resolved(p) for p in excluded_paths}
    filename_filter = (_normalise_filename_text(filename_contains)
                       if filename_contains and filename_contains.strip() else None)
    found: dict[str, Path] = {}
    for folder in folders:
        root = Path(folder)
        if not root.is_dir():
            continue
        iterator = root.rglob("*") if recursive else root.glob("*")
        for path in iterator:
            if (path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES
                    and not path.name.startswith("~$") and _resolved(path) not in excluded
                    and (filename_filter is None
                         or filename_filter in _normalise_filename_text(path.name))):
                found.setdefault(_resolved(path), path.resolve())
    return sorted(found.values(), key=lambda p: str(p).casefold())


def merge_workbooks(paths: Iterable[str | os.PathLike], options: MergeOptions | None = None,
                    base_folders: Iterable[str | os.PathLike] = (),
                    progress_callback=None, is_cancelled=None) -> MergeResult:
    """读取 Excel/CSV/TXT 并按表头合并；单个损坏文件不会中断整个任务。"""
    from openpyxl import load_workbook

    opts = options or MergeOptions()
    if opts.header_row < 1:
        raise ValueError("表头行必须大于或等于 1")
    if opts.column_mode not in {"union", "intersection"}:
        raise ValueError("列合并方式必须是 union 或 intersection")
    if opts.sheet_match_mode not in {"exact", "contains"}:
        raise ValueError("工作表匹配方式必须是 exact 或 contains")
    if opts.formula_mode not in {"values", "formulas"}:
        raise ValueError("公式处理方式必须是 values 或 formulas")
    if opts.duplicate_mode not in {"all", "first", "last"}:
        raise ValueError("重复数据处理方式必须是 all、first 或 last")

    result = MergeResult()
    records: list[dict[str, object]] = []
    display_names: dict[str, str] = {}
    key_orders: list[list[str]] = []
    # 保留该参数用于后续支持相对路径显示；当前始终记录完整目录，避免同名目录混淆。
    _ = base_folders

    paths = list(paths)
    for file_index, raw_path in enumerate(paths, 1):
        _check_cancelled(is_cancelled)
        path = Path(raw_path).resolve()
        if progress_callback:
            progress_callback(file_index - 1, len(paths), path.name)
        if path.suffix.lower() in DELIMITED_SUFFIXES:
            try:
                _extract_delimited(path, opts, result, records, display_names, key_orders,
                                   is_cancelled)
            except MergeCancelled:
                raise
            except Exception as exc:
                result.statuses.append(FileStatus(
                    str(path), "分隔文本", status="失败", message=f"读取失败：{exc}"))
            if progress_callback:
                progress_callback(file_index, len(paths), path.name)
            continue
        wb = None
        formula_wb = None
        try:
            if opts.formula_mode == "formulas":
                wb = load_workbook(path, read_only=True, data_only=False)
            else:
                wb = load_workbook(path, read_only=True, data_only=True)
                # 第二份只用于识别公式及在缓存缺失时保留公式文本。
                formula_wb = load_workbook(path, read_only=True, data_only=False)
        except Exception as exc:
            if wb is not None:
                wb.close()
            if formula_wb is not None:
                formula_wb.close()
            result.statuses.append(FileStatus(str(path), "—", status="失败",
                                              message=f"无法打开：{exc}"))
            continue
        try:
            if opts.sheet_name is None:
                sheet_names = wb.sheetnames
            elif opts.sheet_match_mode == "contains":
                needle = _normalise_sheet_text(opts.sheet_name)
                sheet_names = [name for name in wb.sheetnames
                               if needle in _normalise_sheet_text(name)]
                if not sheet_names:
                    result.statuses.append(FileStatus(
                        str(path), f"名称包含：{opts.sheet_name}", status="跳过",
                        message="没有名称包含指定字符的工作表"))
                    continue
            else:
                sheet_names = [opts.sheet_name]
            for sheet_name in sheet_names:
                if sheet_name not in wb.sheetnames:
                    result.statuses.append(FileStatus(
                        str(path), sheet_name, status="跳过", message="工作表不存在"))
                    continue
                ws = wb[sheet_name]
                formula_ws = formula_wb[sheet_name] if formula_wb is not None else None
                try:
                    _extract_sheet(path, ws, formula_ws, opts, result, records,
                                   display_names, key_orders, is_cancelled)
                except MergeCancelled:
                    raise
                except Exception as exc:
                    result.statuses.append(FileStatus(
                        str(path), sheet_name, status="失败", message=f"读取失败：{exc}"))
        finally:
            wb.close()
            if formula_wb is not None:
                formula_wb.close()
        if progress_callback:
            progress_callback(file_index, len(paths), path.name)

    if not records:
        return result

    if opts.column_mode == "intersection":
        common = set(key_orders[0])
        for keys in key_orders[1:]:
            common.intersection_update(keys)
        data_keys = [key for key in key_orders[0] if key in common]
    else:
        data_keys = list(display_names)

    result.headers = (list(SOURCE_HEADERS) if opts.include_source else []) + [
        display_names[key] for key in data_keys
    ]
    for source in records:
        row = ([source[h] for h in SOURCE_HEADERS] if opts.include_source else [])
        value_map = source["__values__"]
        row.extend(value_map.get(key) for key in data_keys)
        result.rows.append(row)
    sort_keys = _resolve_requested_keys(
        data_keys, display_names, opts.sort_columns, "排序列")
    if opts.sort_by_sample:
        sample_key = _resolve_sample_key(data_keys, display_names, opts.sample_column)
        sort_keys = [sample_key] + [key for key in sort_keys if key != sample_key]

    if opts.duplicate_mode != "all" or opts.duplicate_columns or opts.sort_by_sample:
        duplicate_keys = _resolve_requested_keys(
            data_keys, display_names, opts.duplicate_columns, "重复判定列")
        if not duplicate_keys:
            duplicate_keys = ([sort_keys[0]] if sort_keys else [
                _resolve_sample_key(data_keys, display_names, opts.sample_column)])
        offset = len(SOURCE_HEADERS) if opts.include_source else 0
        indices = [offset + data_keys.index(key) for key in duplicate_keys]
        if opts.duplicate_mode == "all":
            duplicates = _duplicate_count(result.rows, indices)
            result.warnings.append(
                f"重复数据检查列：{'、'.join(display_names[k] for k in duplicate_keys)}；"
                f"发现 {duplicates} 条重复记录，已全部保留")
        else:
            before = len(result.rows)
            result.rows = _deduplicate_rows(result.rows, indices, opts.duplicate_mode)
            removed = before - len(result.rows)
            result.warnings.append(
                f"重复数据检查列：{'、'.join(display_names[k] for k in duplicate_keys)}；"
                f"按“{'保留第一条' if opts.duplicate_mode == 'first' else '保留最后一条'}”"
                f"移除 {removed} 行")

    if sort_keys:
        offset = len(SOURCE_HEADERS) if opts.include_source else 0
        indices = [offset + data_keys.index(key) for key in sort_keys]
        result.rows.sort(key=lambda row: tuple(_natural_sort_key(row[i]) for i in indices))
        result.warnings.append(
            f"合并结果已按“{'、'.join(display_names[k] for k in sort_keys)}”自然排序")
    return result


def _extract_delimited(path, opts, result, records, display_names, key_orders,
                       is_cancelled=None) -> None:
    encoding = _detect_delimited_encoding(path, is_cancelled)
    with path.open("r", encoding=encoding, newline="") as stream:
        sample = stream.read(65536)
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel if path.suffix.lower() == ".csv" else csv.excel_tab
    delimiter = dialect.delimiter
    header_values = None
    max_columns = 0
    # 第一遍只确定表头和最大列数，不把整份文件复制到内存。
    with path.open("r", encoding=encoding, newline="") as stream:
        for row_number, row in enumerate(csv.reader(stream, dialect=dialect), start=1):
            if row_number % 1000 == 0:
                _check_cancelled(is_cancelled)
            if row_number == opts.header_row:
                header_values = row
            if row_number >= opts.header_row:
                max_columns = max(max_columns, len(row))
    if header_values is None:
        result.statuses.append(FileStatus(
            str(path), "分隔文本", status="跳过",
            message=f"文件不足 {opts.header_row} 行，无法读取表头"))
        return
    if not any(not _blank(value) for value in header_values):
        result.statuses.append(FileStatus(
            str(path), "分隔文本", status="跳过",
            message=f"第 {opts.header_row} 行没有表头"))
        return
    raw_headers = list(header_values) + [None] * (max_columns - len(header_values))
    keys, names, header_warnings = _prepare_headers(raw_headers, opts.column_mappings)
    for warning in header_warnings:
        result.warnings.append(f"{path.name} / 分隔文本：{warning}")
    for key, name in zip(keys, names):
        display_names.setdefault(key, name)
    key_orders.append(keys)

    count = 0
    with path.open("r", encoding=encoding, newline="") as stream:
        for row_number, raw in enumerate(csv.reader(stream, dialect=dialect), start=1):
            if row_number <= opts.header_row:
                continue
            if row_number % 1000 == 0:
                _check_cancelled(is_cancelled)
            values = list(raw) + [None] * (max_columns - len(raw))
            if opts.skip_empty_rows and all(_blank(value) for value in values):
                continue
            source = {
                "来源文件夹": str(path.parent),
                "来源文件": path.name,
                "来源工作表": "分隔文本",
                "原始行号": row_number,
                "__values__": {key: value for key, value in zip(keys, values)},
            }
            records.append(source)
            count += 1
    delimiter_name = {",": "逗号", "\t": "制表符", ";": "分号", "|": "竖线"}.get(
        delimiter, repr(delimiter))
    result.statuses.append(FileStatus(
        str(path), "分隔文本", rows=count, columns=len(keys),
        message=f"编码 {encoding}，{delimiter_name}分隔"))


def _detect_delimited_encoding(path: Path, is_cancelled=None) -> str:
    """分块验证编码，避免为了识别编码一次性读取整份文本文件。"""
    for encoding in ("utf-8-sig", "utf-16", "gb18030"):
        try:
            with path.open("r", encoding=encoding, newline="") as stream:
                while stream.read(1024 * 1024):
                    _check_cancelled(is_cancelled)
            return encoding
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise ValueError("无法识别文本编码（支持 UTF-8、UTF-16、GB18030）")


def _resolve_sample_key(data_keys, display_names, requested):
    if requested:
        wanted = _normalise_header(requested)
        for key in data_keys:
            if key == wanted or _normalise_header(display_names[key]) == wanted:
                return key
        raise ValueError(f"未找到指定的样品名列“{requested}”，请检查表头名称")
    aliases = {_normalise_header(name) for name in SAMPLE_NAME_ALIASES}
    for key in data_keys:
        if _normalise_header(display_names[key]) in aliases:
            return key
    raise ValueError(
        "未自动识别到样品名列；请在提取设置中填写实际列名，例如“样品名称”")


def _resolve_requested_keys(data_keys, display_names, requested, label):
    keys = []
    for name in requested:
        wanted = _normalise_header(name)
        match = next((key for key in data_keys
                      if key == wanted or _normalise_header(display_names[key]) == wanted), None)
        if match is None:
            raise ValueError(f"{label}“{name}”不存在，请检查字段名称")
        if match not in keys:
            keys.append(match)
    return keys


def _deduplicate_rows(rows, indices, mode):
    if mode == "first":
        seen, output = set(), []
        for row in rows:
            key = _row_identity(row, indices)
            if key is not None and key in seen:
                continue
            if key is not None:
                seen.add(key)
            output.append(row)
        return output
    seen, output = set(), []
    for row in reversed(rows):
        key = _row_identity(row, indices)
        if key is not None and key in seen:
            continue
        if key is not None:
            seen.add(key)
        output.append(row)
    output.reverse()
    return output


def _duplicate_count(rows, indices):
    seen, duplicates = set(), 0
    for row in rows:
        key = _row_identity(row, indices)
        if key is None:
            continue
        if key in seen:
            duplicates += 1
        else:
            seen.add(key)
    return duplicates


def _row_identity(row, indices):
    values = tuple(row[index] for index in indices)
    if all(_blank(value) for value in values):
        return None
    return tuple(_normalise_header(str(value)) if isinstance(value, str) else value
                 for value in values)


def _natural_sort_key(value):
    if _blank(value):
        return (1, ())
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return (0, ((0, float(value)),))
    text = unicodedata.normalize("NFKC", str(value)).strip().casefold()
    tokens = tuple((0, int(part)) if part.isdigit() else (1, part)
                   for part in re.split(r"(\d+)", text) if part)
    return (0, tokens)


def _extract_sheet(path, ws, formula_ws, opts, result, records,
                   display_names, key_orders, is_cancelled=None) -> None:
    header_cells = next(ws.iter_rows(min_row=opts.header_row, max_row=opts.header_row), None)
    formula_header_cells = (next(formula_ws.iter_rows(
        min_row=opts.header_row, max_row=opts.header_row), None)
        if formula_ws is not None else None)
    header_values, missing_formula_count = _effective_row(
        header_cells, formula_header_cells, opts.formula_mode)
    if not header_values or not any(not _blank(v) for v in header_values):
        result.statuses.append(FileStatus(str(path), ws.title, status="跳过",
                                          message=f"第 {opts.header_row} 行没有表头"))
        return

    last = max(i for i, value in enumerate(header_values) if not _blank(value)) + 1
    raw_headers = list(header_values[:last])
    keys, names, header_warnings = _prepare_headers(raw_headers, opts.column_mappings)
    for warning in header_warnings:
        result.warnings.append(f"{path.name} / {ws.title}：{warning}")
    for key, name in zip(keys, names):
        display_names.setdefault(key, name)
    key_orders.append(keys)

    count = 0
    folder_label = str(path.parent)
    value_rows = ws.iter_rows(min_row=opts.header_row + 1, max_col=last)
    formula_rows = (formula_ws.iter_rows(min_row=opts.header_row + 1, max_col=last)
                    if formula_ws is not None else None)
    for row_number, value_cells in enumerate(value_rows, start=opts.header_row + 1):
        if row_number % 1000 == 0:
            _check_cancelled(is_cancelled)
        formula_cells = next(formula_rows, None) if formula_rows is not None else None
        values, missing = _effective_row(value_cells, formula_cells, opts.formula_mode)
        missing_formula_count += missing
        if opts.skip_empty_rows and all(_blank(v) for v in values):
            continue
        value_map = {key: value for key, value in zip(keys, values)}
        source = {
            "来源文件夹": folder_label,
            "来源文件": path.name,
            "来源工作表": ws.title,
            "原始行号": row_number,
            "__values__": value_map,
        }
        records.append(source)
        count += 1
    result.statuses.append(FileStatus(str(path), ws.title, rows=count, columns=len(keys),
                                      message="无数据行" if count == 0 else ""))
    if missing_formula_count:
        result.warnings.append(
            f"{path.name} / {ws.title}：{missing_formula_count} 个公式单元格没有缓存结果，"
            "已保留公式文本，请用 Excel 重新计算源文件后再提取以获得数值")


def _effective_row(value_cells, formula_cells, formula_mode):
    """返回一行有效值；缓存缺失时不静默丢弃公式。"""
    if value_cells is None:
        return [], 0
    if formula_mode == "formulas" or formula_cells is None:
        return [cell.value for cell in value_cells], 0
    values, missing = [], 0
    for value_cell, formula_cell in zip(value_cells, formula_cells):
        value = value_cell.value
        if formula_cell.data_type == "f" and value is None:
            value = formula_cell.value
            missing += 1
        values.append(value)
    return values, missing


def _prepare_headers(values: list[object], mappings=()) -> tuple[list[str], list[str], list[str]]:
    keys, names, warnings = [], [], []
    used: dict[str, int] = {}
    mapping = {_normalise_header(source): target.strip()
               for source, target in mappings if source.strip() and target.strip()}
    for index, value in enumerate(values, 1):
        name = str(value).strip() if not _blank(value) else f"未命名列_{index}"
        if _blank(value):
            warnings.append(f"第 {index} 列表头为空，已命名为“{name}”")
        mapped = mapping.get(_normalise_header(name))
        if mapped:
            warnings.append(f"字段“{name}”已统一为“{mapped}”")
            name = mapped
        base = _normalise_header(name)
        number = used.get(base, 0) + 1
        used[base] = number
        key = base if number == 1 else f"{base}__{number}"
        if number > 1:
            renamed = f"{name}_{number}"
            warnings.append(f"表头“{name}”重复，已命名为“{renamed}”")
            name = renamed
        keys.append(key)
        names.append(name)
    return keys, names, warnings


def export_merged_workbook(path: str | os.PathLike, result: MergeResult,
                           max_rows_per_sheet: int = EXCEL_MAX_ROWS) -> None:
    """原子写出合并结果和处理日志，避免中途失败留下不完整文件。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from ..fileio.excel_safe import append_safe_row

    target = Path(path)
    if target.suffix.lower() != ".xlsx":
        target = target.with_suffix(".xlsx")
    if not result.headers or not result.rows:
        raise ValueError("没有可导出的合并数据")
    if max_rows_per_sheet < 2 or max_rows_per_sheet > EXCEL_MAX_ROWS:
        raise ValueError(f"每个工作表行数必须在 2 到 {EXCEL_MAX_ROWS:,} 之间")
    target.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook(write_only=False)
    fill = PatternFill("solid", fgColor="1F4E78")
    data_capacity = max_rows_per_sheet - 1
    sheet_count = (len(result.rows) + data_capacity - 1) // data_capacity
    for sheet_index in range(sheet_count):
        ws = wb.active if sheet_index == 0 else wb.create_sheet()
        ws.title = "合并数据" if sheet_index == 0 else f"合并数据_{sheet_index + 1}"
        append_safe_row(ws, result.headers)
        start = sheet_index * data_capacity
        for row in result.rows[start:start + data_capacity]:
            append_safe_row(ws, row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center")
        for idx, header in enumerate(result.headers, 1):
            ws.column_dimensions[get_column_letter(idx)].width = min(
                max(len(header) + 2, 12), 32)

    log = wb.create_sheet("合并日志")
    append_safe_row(log, ["状态", "文件路径", "工作表", "提取行数", "列数", "说明"])
    for status in result.statuses:
        append_safe_row(log, [status.status, status.path, status.sheet, status.rows,
                              status.columns, status.message])
    for warning in result.warnings:
        append_safe_row(log, ["警告", "", "", "", "", warning])
    if sheet_count > 1:
        append_safe_row(log, ["信息", "", "", "", "",
                              f"数据超过单表容量，已拆分为 {sheet_count} 个数据工作表"])
    log.freeze_panes = "A2"
    log.auto_filter.ref = log.dimensions
    for cell in log[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill

    fd, temp_name = tempfile.mkstemp(prefix="excel_merge_", suffix=".xlsx",
                                     dir=str(target.parent))
    os.close(fd)
    try:
        wb.save(temp_name)
        os.replace(temp_name, target)
    finally:
        wb.close()
        if os.path.exists(temp_name):
            os.unlink(temp_name)


def _normalise_header(value: str) -> str:
    return "".join(unicodedata.normalize("NFKC", value).strip().casefold().split())


def _normalise_sheet_text(value: str) -> str:
    """工作表匹配忽略大小写和全半角，但保留名称内部空格。"""
    return unicodedata.normalize("NFKC", value).strip().casefold()


def _normalise_filename_text(value: str) -> str:
    """文件名筛选忽略大小写、全半角和首尾空格。"""
    return unicodedata.normalize("NFKC", value).strip().casefold()


def _blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _check_cancelled(is_cancelled):
    if is_cancelled is not None and is_cancelled():
        raise MergeCancelled("用户已取消提取")


def _resolved(path) -> str:
    return os.path.normcase(os.path.abspath(os.fspath(path)))
