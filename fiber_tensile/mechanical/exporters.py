"""Recompute on export; preserve parameters, provenance and unrounded data."""

import csv
from io import BytesIO
import json
import os
from pathlib import Path

from ..config import APP_VERSION
from ..fileio.atomic import atomic_write, atomic_write_bytes
from ..fileio.excel_safe import append_safe_row, spreadsheet_safe_text
from .core import ALGORITHM_VERSION, MODES, SOURCES, analyze, summarize

HEADERS = ["样品", "试验", "变形来源", "模量(GPa)", "峰值应力(MPa)", "峰值应变(%)",
           "末点应变(%)", "指定应变(%)", "指定应变应力(MPa)", "首点至峰值积分(MJ/m³)",
           "R²", "人工断裂点号", "断裂应变(%)", "断裂点应力(MPa)",
           "测试日期", "备注", "来源文件", "软件版本", "算法版本", "提示"]


def recalculated(entries):
    if not entries:
        raise ValueError("请选择至少一条分析结果")
    return [{**e, "result": analyze(e["x"], e["force"], e["spec"])} for e in entries]


def summary_row(e):
    r, s = e["result"], e["spec"]
    return [e["name"], MODES[s.mode], SOURCES[s.source], r.modulus_gpa, r.strength_mpa,
            r.peak_strain_pct, r.end_strain_pct, s.target_pct, r.target_stress_mpa,
            r.work_to_peak_mj_m3, r.r2, s.fracture_point or None, r.fracture_strain_pct,
            r.fracture_stress_mpa, e.get("test_date", ""), e.get("remark", ""),
            e.get("import", {}).get("file", ""), APP_VERSION, ALGORITHM_VERSION,
            "；".join(r.warnings)]


def grouped_statistics(entries):
    groups = {}
    for e in entries:
        s = e["spec"]
        key = (s.mode, s.source, s.fit_mode, s.fit_lo, s.fit_hi)
        groups.setdefault(key, []).append(e["result"])
    return [(key, summarize(results)) for key, results in groups.items()]


def parameter_rows(s):
    rows = [("试验", MODES[s.mode]), ("变形来源", SOURCES[s.source]),
            ("截面", {"rectangle": "矩形", "circle": "圆形", "area": "直接面积"}[s.section])]
    fields = {"rectangle": [("宽度(mm)", s.width_mm), ("厚度(mm)", s.thickness_mm)],
              "circle": [("直径(mm)", s.diameter_mm)], "area": [("截面积(mm²)", s.area_mm2)]}
    rows.extend(fields[s.section])
    if s.source != "strain":
        rows.append(("标距(mm)", s.gauge_mm))
    rows.extend([("拟合区间依据", "工程应变(%)" if s.fit_mode == "strain" else "峰值力比例(%)"),
                 ("拟合下限", s.fit_lo), ("拟合上限", s.fit_hi), ("指定应变(%)", s.target_pct),
                 ("人工确认断裂点号", s.fracture_point or "未指定")])
    return rows


def export_csv(path, entries):
    entries = recalculated(entries)
    def write(temporary):
        with open(temporary, "w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(HEADERS)
            for e in entries:
                writer.writerow([spreadsheet_safe_text(v) for v in summary_row(e)])
    atomic_write(path, write)


def export_excel(path, entries):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Font, PatternFill
    from .plotting import build_figure

    entries = recalculated(entries)
    book = Workbook()
    summary = book.active
    summary.title = "结果汇总"
    append_safe_row(summary, HEADERS)
    for e in entries:
        append_safe_row(summary, summary_row(e))
    stats = book.create_sheet("分组统计")
    append_safe_row(stats, ["试验", "来源", "拟合方式", "下限", "上限", "指标", "数量", "均值", "样本标准差", "CV(%)"])
    labels = {"modulus_gpa": "模量(GPa)", "strength_mpa": "峰值应力(MPa)", "peak_strain_pct": "峰值应变(%)"}
    for key, values in grouped_statistics(entries):
        for name, v in values.items():
            append_safe_row(stats, [MODES[key[0]], SOURCES[key[1]], *key[2:], labels[name], v["n"], v["mean"], v["sd"], v["cv_pct"]])
    metadata = book.create_sheet("参数与过程")
    append_safe_row(metadata, ["样品", "项目", "内容"])
    for e in entries:
        for key, value in parameter_rows(e["spec"]):
            append_safe_row(metadata, [e["name"], key, value])
        for key, value in {"导入设置": e.get("import", {}),
                           "保存时版本": [e.get("saved_app_version", APP_VERSION), e.get("saved_algorithm_version", ALGORITHM_VERSION)]}.items():
            append_safe_row(metadata, [e["name"], key, json.dumps(value, ensure_ascii=False)])
        for line in e["result"].steps + e["result"].warnings:
            append_safe_row(metadata, [e["name"], "计算过程 / 提示", line])
    for i, e in enumerate(entries, 1):
        # Split at Excel's row limit without dropping points.
        page = None
        for j, (x, f, strain, stress) in enumerate(zip(e["x"], e["force"], e["result"].strain, e["result"].stress_mpa)):
            if j % 1_048_575 == 0:
                page = book.create_sheet(f"曲线{i}-{j // 1_048_575 + 1}")
                append_safe_row(page, ["点号", "原始应变(mm/mm)" if e["spec"].source == "strain" else "原始变形(mm)", "力(N)", "工程应变(%)", "工程应力(MPa)"])
            append_safe_row(page, [j + 1, x, f, strain * 100, stress])
    chart = book.create_sheet("曲线对比")
    chart["A1"] = "图显示前20条，每条最多约5000点；完整曲线数据见各曲线工作表。"
    buf = BytesIO()
    fig = build_figure(entries)
    fig.savefig(buf, format="png", dpi=120)
    buf.seek(0)
    chart.add_image(Image(buf), "A3")
    for sheet in book:
        sheet.freeze_panes = "A2"
        if sheet != chart:
            sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E79")
        for col in ("A", "B", "C", "D", "E"):
            sheet.column_dimensions[col].width = 24
    metadata.column_dimensions["C"].width = 100
    atomic_write(path, book.save)
    book.close()


def export_pdf(path, entries):
    import html
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont, TTFError
    from reportlab.platypus import SimpleDocTemplate, Paragraph, PageBreak, Image
    from .plotting import build_figure

    entries = recalculated(entries)
    font = "MechanicalCJK"
    if font not in pdfmetrics.getRegisteredFontNames():
        # Embed a subset of a Windows CJK font: CID viewer substitution loses
        # superscripts/minus signs in formulae on otherwise valid PDFs.
        for filename in ("msyh.ttc", "simsun.ttc"):
            candidate = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts" / filename
            if candidate.is_file():
                try:
                    pdfmetrics.registerFont(TTFont(font, str(candidate), subfontIndex=0))
                    break
                except TTFError:
                    continue
        else:
            raise ValueError("PDF 导出需要可嵌入的微软雅黑或宋体字体，请检查 Windows 字体安装。")
    body = ParagraphStyle("body", fontName=font, fontSize=10, leading=15, spaceAfter=5, wordWrap="CJK")
    title = ParagraphStyle("title", parent=body, fontSize=16, leading=22, spaceAfter=12)
    p = lambda value, style=body: Paragraph(html.escape(str(value)), style)
    story = [p("通用力学曲线分析报告", title), p(f"重算版本：{APP_VERSION} / {ALGORITHM_VERSION}")]
    for i, e in enumerate(entries):
        if i:
            story.append(PageBreak())
        story.extend([p(e["name"], title), p(f"测试日期：{e.get('test_date', '')}；备注：{e.get('remark', '')}"),
                      p(f"来源：{e.get('import', {}).get('file', '')}"),
                      p("；".join(f"{key}：{value}" for key, value in parameter_rows(e["spec"])))])
        for line in e["result"].steps + e["result"].warnings:
            story.append(p(line))
        buf = BytesIO()
        build_figure([e]).savefig(buf, format="png", dpi=100)
        buf.seek(0)
        story.append(Image(buf, width=450, height=253))
    output = BytesIO()
    SimpleDocTemplate(output, leftMargin=42, rightMargin=42, topMargin=36, bottomMargin=36).build(story)
    atomic_write_bytes(path, output.getvalue())
