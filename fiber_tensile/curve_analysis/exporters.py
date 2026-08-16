"""单丝曲线分析报告的 Excel/PDF 导出。"""

from datetime import datetime


def _build_rows(sample_name, diameter_um, gauge_mm, result, remark, test_date):
    info = [("样品名称", sample_name), ("直径(μm)", f"{diameter_um:g}"),
            ("标距(mm)", f"{gauge_mm:g}"), ("测试日期", test_date or "—"),
            ("备注", remark or "—")]
    results = [("弹性模量 E(GPa)", f"{result.e_gpa:.2f}"),
               ("断裂强度 σ_b(MPa)", f"{result.strength_mpa:.2f}"),
               ("断裂伸长率（校正后, %）", f"{result.elongation_pct:.3f}"),
               ("平台段应变（已校正, %）", f"{result.toe_eps_pct:.3f}"),
               ("拟合优度 R²", f"{result.r2:.4f}"),
               ("数据点数", str(result.n_points))]
    return info, results, result.steps


def export_excel_report(path, sample_name, diameter_um, gauge_mm, result,
                        disp_mm, force_n, remark=None, test_date=None) -> None:
    from io import BytesIO

    import matplotlib.pyplot as plt
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    from .plotting import build_curve_figure

    info, results, steps = _build_rows(sample_name, diameter_um, gauge_mm,
                                       result, remark, test_date)
    wb = Workbook()
    ws = wb.active
    ws.title = "曲线分析"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 24

    row = 1
    c = ws.cell(row, 1, f"单丝拉伸曲线分析报告 — {sample_name}")
    c.font = Font(bold=True, size=14)
    row += 2
    for section, pairs in (("样品信息", info), ("计算结果", results)):
        c = ws.cell(row, 1, section)
        c.font = Font(bold=True, size=12)
        row += 1
        for k, v in pairs:
            ws.cell(row, 1, k).font = Font(bold=True)
            ws.cell(row, 2, v)
            row += 1
        row += 1
    c = ws.cell(row, 1, "计算过程")
    c.font = Font(bold=True, size=12)
    row += 1
    for line in steps:
        ws.cell(row, 1, line)
        row += 1
    fig = build_curve_figure(disp_mm, force_n, result)
    buf = BytesIO()
    fig.savefig(buf, format="png", dpi=150)
    plt.close(fig)
    img = XLImage(buf)
    img.width = 640
    img.height = 460
    ws.add_image(img, "D2")
    wb.save(path)


def export_pdf_report(path, sample_name, diameter_um, gauge_mm, result,
                      disp_mm, force_n, remark=None, test_date=None) -> None:
    from io import BytesIO

    import matplotlib.pyplot as plt
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import (Image, Paragraph, SimpleDocTemplate, Spacer,
                                    Table, TableStyle)

    from .plotting import build_curve_figure

    cjk = "STSong-Light"
    if cjk not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont(cjk))
    title = ParagraphStyle("t", fontName=cjk, fontSize=16, leading=22,
                           alignment=1, spaceAfter=14)
    h2 = ParagraphStyle("h", fontName=cjk, fontSize=13, leading=18,
                        spaceBefore=14, spaceAfter=8,
                        textColor=colors.HexColor("#1F4E79"))
    body = ParagraphStyle("b", fontName=cjk, fontSize=10, leading=14)
    grid = TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B0B0B0")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ])

    def p(text, style=body):
        import html as _html
        return Paragraph(_html.escape(str(text)), style)

    info, results, steps = _build_rows(sample_name, diameter_um, gauge_mm,
                                       result, remark, test_date)
    doc = SimpleDocTemplate(str(path), pagesize=A4,
                            leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm,
                            title=f"单丝拉伸曲线分析报告 - {sample_name}")
    story = [p("单丝拉伸曲线分析报告", title),
             p(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}",
               ParagraphStyle("meta", parent=body, alignment=1, fontSize=9)),
             p("一、样品信息", h2),
             Table([[p(k), p(v)] for k, v in info],
                   colWidths=[4 * cm, 11 * cm], style=grid),
             p("二、计算结果", h2),
             Table([[p(k), p(v)] for k, v in results],
                   colWidths=[6 * cm, 9 * cm], style=grid)]
    fig = build_curve_figure(disp_mm, force_n, result)
    buf = BytesIO()
    fig.savefig(buf, format="png", dpi=150)
    plt.close(fig)
    story.append(Spacer(1, 0.3 * cm))
    story.append(Image(buf, width=15 * cm, height=10.7 * cm))
    story.append(p("三、计算过程", h2))
    for line in steps:
        story.append(p(line))
    doc.build(story)
