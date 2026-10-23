"""Opt-in packaged verification with synthetic inputs and isolated user state."""

import json
from pathlib import Path


def run(output_path):
    from scripts.test_environment import isolated_environment

    output = Path(output_path).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    with isolated_environment() as root:
        import logging
        import numpy as np
        from PySide6.QtWidgets import QApplication
        from .config import APP_VERSION, default_db_path
        from .data.db import open_database
        from .ui.main_window import MainWindow
        from .ui.prefs import _settings
        from .mechanical.core import Spec, analyze
        from .mechanical.exporters import export_excel, export_pdf, export_csv
        from .mechanical.repository import save_entries, load_entry
        from .core.models import Batch, Sample, MODE_FILAMENT
        from .data.repository import create_batch, add_samples, get_batch, get_strengths
        from .core.statistics import batch_statistics
        from .core.weibull import weibull_fit
        from .fileio.exporter_excel import export_excel_report
        from .fileio.exporter_pdf import export_pdf_report
        from openpyxl import load_workbook

        app = QApplication([])
        conn = open_database(default_db_path())
        win = None
        try:
            assert Path(_settings().fileName()).resolve().is_relative_to(root)
            win = MainWindow(conn)
            win.show()
            win._nav.setCurrentRow(5)
            app.processEvents()
            assert win._nav.count() == 6
            x = np.linspace(0, .02, 201).tolist()
            force = [v * 60000 for v in x]
            spec = Spec(fracture_point=201)
            result = analyze(x, force, spec)
            assert abs(result.modulus_gpa - 3) < 1e-9
            entry = dict(name="打包合成验证", x=x, force=force, spec=spec, result=result)
            record_id = save_entries(conn, [entry])[0]
            entry = load_entry(conn, record_id)
            win.mechanical_tab._add_entry(entry)
            win.mechanical_tab._table.selectRow(0)
            app.processEvents()
            for extension, exporter in (("xlsx", export_excel), ("pdf", export_pdf), ("csv", export_csv)):
                path = root / f"report.{extension}"
                exporter(path, [entry])
                assert path.stat().st_size > 100
            # 确认 EXE 中实际包含本次 GPa 横坐标及 Origin 拟合线导出修复。
            batch_id = create_batch(conn, Batch(batch_name="Weibull 合成验证", mode=MODE_FILAMENT))
            add_samples(conn, batch_id, [
                Sample(breaking_force=force, diameter_um=7.0)
                for force in (7.1, 8.3, 9.5, 10.2, 12.4)
            ])
            batch = get_batch(conn, batch_id)
            strengths = get_strengths(conn, batch_id)
            win._nav.setCurrentRow(0)
            win.analysis_tab.refresh()
            win.analysis_tab._combo.setCurrentIndex(win.analysis_tab._combo.findData(batch_id))
            win.analysis_tab._combo_wfit.setCurrentIndex(0)
            win.analysis_tab._combo_rank.setCurrentIndex(1)
            app.processEvents()
            ax = win.analysis_tab._canvas.fig.axes[0]
            assert "GPa" in ax.get_xlabel()
            fit = weibull_fit(strengths, "mean_rank")
            stats = batch_statistics(strengths)
            export_excel_report(root / "weibull.xlsx", batch, stats, fit)
            export_pdf_report(root / "weibull.pdf", batch, stats, fit)
            workbook = load_workbook(root / "weibull.xlsx")
            try:
                scatter = list(workbook["Weibull 分析数据"].values)
                assert scatter[0][1] == "σ(GPa)" and scatter[0][3] == "ln(σ/GPa)"
                np.testing.assert_allclose([r[1] for r in scatter[1:]], np.sort(strengths) / 1000)
                np.testing.assert_allclose([r[3:5] for r in scatter[1:]], ax.collections[0].get_offsets())
                line = list(workbook["Weibull 拟合直线数据"].iter_rows(
                    min_row=2, max_row=101, min_col=1, max_col=2, values_only=True))
                assert len(line) == len(ax.lines[0].get_xdata()) == 100
                np.testing.assert_allclose(np.asarray(line)[:, 0], ax.lines[0].get_xdata())
                np.testing.assert_allclose(np.asarray(line)[:, 1], ax.lines[0].get_ydata())
            finally:
                workbook.close()
            checks = dict(version=APP_VERSION, navigation=6, modulus_gpa=result.modulus_gpa,
                          database_isolated=True, settings_isolated=True,
                          exports=["xlsx", "pdf", "csv"], history=True,
                          weibull_gpa=True, weibull_line_points=100)
            output.write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
        finally:
            if win is not None:
                win.close()
            conn.close()
            logging.shutdown()
    return 0
