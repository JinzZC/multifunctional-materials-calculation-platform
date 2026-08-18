"""复合材料基础计算报告的 Excel/PDF 导出。"""

from datetime import datetime
from ..fileio.excel_safe import set_cell_value

_TYPE_LABELS = {"rom": "单层板工程常数（混合法则）", "volume_fraction": "纤维体积分数/孔隙率"}


def _build_sections(calc_type, sample_name, inputs, result_steps, result_pairs,
                    remark, test_date):
    info = [("样品名称", sample_name), ("计算类型", _TYPE_LABELS[calc_type]),
            ("测试日期", test_date or "—"), ("备注", remark or "—")]
    input_rows = [(k, f"{v:g}" if isinstance(v, float) else str(v))
                  for k, v in inputs.items()]
    return info, input_rows, result_pairs, result_steps


def export_excel_report(path, calc_type, sample_name, inputs,
                        result_pairs, steps, remark=None, test_date=None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    info, input_rows, result_pairs, steps = _build_sections(
        calc_type, sample_name, inputs, steps, result_pairs, remark, test_date)
    wb = Workbook()
    ws = wb.active
    ws.title = "复合材料计算"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 26

    row = 1
    c = ws.cell(row, 1, f"复合材料基础计算报告 — {sample_name}")
    c.font = Font(bold=True, size=14)
    row += 2
    for section, pairs in (("样品信息", info), ("输入参数", input_rows),
                           ("计算结果", result_pairs)):
        c = ws.cell(row, 1, section)
        c.font = Font(bold=True, size=12)
        row += 1
        for k, v in pairs:
            ws.cell(row, 1, k).font = Font(bold=True)
            set_cell_value(ws.cell(row, 2), v)
            row += 1
        row += 1
    c = ws.cell(row, 1, "计算过程")
    c.font = Font(bold=True, size=12)
    row += 1
    for line in steps:
        set_cell_value(ws.cell(row, 1), line)
        row += 1
    wb.save(path)


def export_pdf_report(path, calc_type, sample_name, inputs,
                      result_pairs, steps, remark=None, test_date=None) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

    cjk = "STSong-Light"
    if cjk not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont(cjk))
    title = ParagraphStyle("t", fontName=cjk, fontSize=16, leading=22,
                           alignment=1, spaceAfter=14)
    h2 = ParagraphStyle("h", fontName=cjk, fontSize=13, leading=18,
                        spaceBefore=14, spaceAfter=8,
                        textColor=colors.HexColor("#1F4E79"))
    body = ParagraphStyle("b", fontName=cjk, fontSize=10, leading=14)
    grid = TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B0B0B0")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ])

    def p(text, style=body):
        import html as _html
        return Paragraph(_html.escape(str(text)), style)

    info, input_rows, result_pairs, steps = _build_sections(
        calc_type, sample_name, inputs, steps, result_pairs, remark, test_date)
    doc = SimpleDocTemplate(str(path), pagesize=A4,
                            leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm,
                            title=f"复合材料基础计算报告 - {sample_name}")
    story = [p("复合材料基础计算报告", title),
             p(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}",
               ParagraphStyle("meta", parent=body, alignment=1, fontSize=9))]
    for sec_name, pairs in (("一、样品信息", info), ("二、输入参数", input_rows),
                            ("三、计算结果", result_pairs)):
        story.append(p(sec_name, h2))
        story.append(Table([[p(k), p(v)] for k, v in pairs],
                           colWidths=[6 * cm, 9 * cm], style=grid))
    story.append(p("四、计算过程", h2))
    for line in steps:
        story.append(p(line))
    doc.build(story)
