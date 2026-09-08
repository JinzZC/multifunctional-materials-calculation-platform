from dataclasses import asdict
import json
import sqlite3

import numpy as np
import pytest

from fiber_tensile.data.db import open_database, MIGRATIONS
from fiber_tensile.data.search import search_all
from fiber_tensile.mechanical.core import Spec, analyze
from fiber_tensile.mechanical.importer import read_curve
from fiber_tensile.mechanical.repository import save_entries, list_entries, load_entry, delete_entries
from fiber_tensile.mechanical.exporters import export_excel, export_csv, export_pdf, grouped_statistics


def entry(name="试样A"):
    x = np.linspace(0, .02, 101).tolist()
    force = [v * 60000 for v in x]
    spec = Spec()
    return {"name": name, "x": x, "force": force, "spec": spec,
            "result": analyze(x, force, spec), "import": {"file": "instrument.csv"},
            "test_date": "2026-09-08", "remark": "合成数据"}


def test_csv_units_and_automatic_header(tmp_path):
    path = tmp_path / "curve.csv"
    path.write_text("strain(%),force(kN)\n" + "\n".join(f"{i / 10},{i / 100}" for i in range(21)), encoding="utf-8")
    x, force, origin = read_curve(path, source="strain", x_unit="%", force_unit="kN")
    assert x[-1] == pytest.approx(.02)
    assert force[-1] == pytest.approx(200)
    assert origin["column_mapping"] == {"x": "strain(%)", "force": "force(kN)"}
    with pytest.raises(ValueError, match="不一致"):
        read_curve(path, source="strain", x_unit="%", force_unit="N")


def test_custom_workbook_sheet_and_mapping(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / "instrument.xlsx"
    book = Workbook()
    sheet = book.create_sheet("曲线")
    sheet.append(["仪器导出"])
    sheet.append(["通道A", "通道B"])
    for i in range(10):
        sheet.append([i * 10, i * 100])
    book.save(path)
    x, f, origin = read_curve(path, source="extension", x_unit="μm", force_unit="mN",
                              sheet_name="曲线", header_row=2, column_mapping={"x": "通道A", "force": "通道B"})
    assert x[-1] == pytest.approx(.09)
    assert f[-1] == pytest.approx(.9)
    assert origin["header_row"] == 2


def test_high_precision_decimal_comma_is_never_treated_as_thousands(tmp_path):
    path = tmp_path / "europe.tsv"
    path.write_text("strain\tforce\n" + "\n".join(f"0,00{i}\t{i}" for i in range(6)), encoding="utf-8")
    with pytest.raises(ValueError, match="分隔符"):
        read_curve(path, source="strain", x_unit="mm/mm", force_unit="N")
    x, _, _ = read_curve(path, source="strain", x_unit="mm/mm", force_unit="N", decimal_separator=",")
    assert x[-1] == pytest.approx(.005)


@pytest.mark.parametrize("content, message", [
    ("strain,force\n0,0\n0.1,NaN", "第 3 行"),
    ("strain,force\n0,0\n0.1,inf", "第 3 行"),
    ("strain,extension,force\n0,0,0", "多个候选"),
])
def test_bad_file_rejected_without_silent_row_drops(tmp_path, content, message):
    path = tmp_path / "bad.csv"
    path.write_text(content, encoding="utf-8")
    with pytest.raises(ValueError, match=message):
        read_curve(path, source="strain", x_unit="mm/mm", force_unit="N")


def test_migration_preserves_v3_data_and_backup(tmp_path):
    path = tmp_path / "old.db"
    raw = sqlite3.connect(path)
    raw.execute("CREATE TABLE schema_migrations (version INTEGER PRIMARY KEY)")
    for version, sql in MIGRATIONS[:3]:
        raw.executescript(sql)
        raw.execute("INSERT INTO schema_migrations VALUES (?)", (version,))
        raw.commit()
    raw.execute("INSERT INTO batches (batch_name,mode) VALUES ('既有样品','filament')")
    raw.commit()
    raw.close()
    conn = open_database(str(path))
    try:
        assert conn.execute("SELECT batch_name FROM batches").fetchone()[0] == "既有样品"
        assert conn.execute("SELECT COUNT(*) FROM mechanical_curves").fetchone()[0] == 0
        backups = list((tmp_path / "backups").glob("*.db"))
        assert len(backups) == 1
        with sqlite3.connect(backups[0]) as snapshot:
            assert snapshot.execute("SELECT MAX(version) FROM schema_migrations").fetchone()[0] == 3
    finally:
        conn.close()


def test_save_load_search_delete_and_transaction_rollback(tmp_path):
    conn = open_database(str(tmp_path / "history.db"))
    try:
        data = entry()
        record_id = save_entries(conn, [data])[0]
        restored = load_entry(conn, record_id)
        assert restored["x"] == data["x"]
        assert asdict(restored["spec"]) == asdict(data["spec"])
        assert restored["result"].modulus_gpa == pytest.approx(3)
        rows = list_entries(conn)
        assert len(rows) == 1
        assert "strain" not in json.loads(rows[0]["result_json"])
        assert search_all(conn, "试样A")[0].module_index == 5
        with pytest.raises(ValueError):
            save_entries(conn, [entry("first"), entry(" ")])
        assert len(list_entries(conn)) == 1
        delete_entries(conn, [record_id])
        assert not list_entries(conn)
    finally:
        conn.close()


def test_exports_recalculate_and_preserve_data(tmp_path):
    from openpyxl import load_workbook
    data = entry("=HYPERLINK(\"bad\")")
    data["result"].modulus_gpa = -123  # stale cache must never be exported
    xlsx, csv, pdf = [tmp_path / f"report.{ext}" for ext in ("xlsx", "csv", "pdf")]
    export_excel(xlsx, [data])
    book = load_workbook(xlsx)
    try:
        assert book["结果汇总"]["A2"].data_type == "s"
        assert book["结果汇总"]["D2"].value == pytest.approx(3)
        assert book["曲线1-1"].max_row == 102
        assert book["曲线1-1"]["E102"].value == pytest.approx(60)
        assert len(book["曲线对比"]._images) == 1
    finally:
        book.close()
    export_csv(csv, [data])
    assert "'=HYPERLINK" in csv.read_text(encoding="utf-8-sig")
    export_pdf(pdf, [data])
    assert pdf.read_bytes().startswith(b"%PDF-")


def test_statistics_keep_apparent_and_measured_results_separate():
    data = entry()
    apparent = entry("crosshead")
    apparent["spec"] = Spec(source="crosshead")
    assert len(grouped_statistics([data, apparent])) == 2


def test_legacy_tensile_snapshot_remains_readable(tmp_path):
    conn = open_database(str(tmp_path / "legacy-mechanical.db"))
    try:
        record_id = save_entries(conn, [entry()])[0]
        parameters = asdict(Spec())
        parameters.update(span_mm=64, load_span_mm=32)
        conn.execute("UPDATE mechanical_curves SET spec_json=? WHERE id=?",
                     (json.dumps(parameters), record_id))
        conn.commit()
        assert load_entry(conn, record_id)["result"].modulus_gpa == pytest.approx(3)
        conn.execute("UPDATE mechanical_curves SET mode='bend4' WHERE id=?", (record_id,))
        conn.commit()
        assert not list_entries(conn)
        assert not search_all(conn, "试样A")
        with pytest.raises(ValueError, match="仅支持拉伸"):
            load_entry(conn, record_id)
        assert conn.execute("SELECT COUNT(*) FROM mechanical_curves").fetchone()[0] == 1
    finally:
        conn.close()
