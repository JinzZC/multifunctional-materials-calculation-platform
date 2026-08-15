"""surface_energy 导入与报告导出测试。"""

import math

import openpyxl
import pytest

from fiber_tensile.fileio.importer import ImportFileError
from fiber_tensile.surface_energy.core import METHOD_OWRK, METHOD_ZISMAN, LiquidPoint
from fiber_tensile.surface_energy.exporters import export_excel_report, export_pdf_report
from fiber_tensile.surface_energy.importer import parse_measurement_file

HEADERS = ["样品名称", "方法", "液体名称", "γ_L(mN/m)", "γ^d(mN/m)",
           "γ^p(mN/m)", "接触角(°)", "备注", "测试日期"]
ROWS = [
    ["PTFE 样品", "双液法", "水", 72.8, 21.8, 51.0, 108.0, "测试", "2026-08-15"],
    ["PTFE 样品", "双液法", "二碘甲烷", 50.8, 50.8, 0.0, 77.0, None, None],
]


def write_xlsx(path, headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    wb.save(path)


def write_csv(path, headers, rows, encoding="utf-8-sig"):
    with open(path, "w", encoding=encoding, newline="") as f:
        f.write(",".join(headers) + "\n")
        for r in rows:
            f.write(",".join("" if v is None else str(v) for v in r) + "\n")


def test_parse_xlsx_basic(tmp_path):
    p = tmp_path / "se.xlsx"
    write_xlsx(p, HEADERS, ROWS)
    preview = parse_measurement_file(p)
    assert preview.sample_name == "PTFE 样品"
    assert preview.method == METHOD_OWRK
    assert preview.remark == "测试"
    assert preview.test_date == "2026-08-15"
    assert len(preview.points) == 2
    assert preview.points[0].liquid_name == "水"
    assert preview.points[0].gamma_d == 21.8
    assert preview.errors == []


def test_parse_csv_gbk(tmp_path):
    p = tmp_path / "se.csv"
    write_csv(p, HEADERS, ROWS, encoding="gb18030")
    preview = parse_measurement_file(p)
    assert len(preview.points) == 2
    assert preview.errors == []


def test_parse_error_rows(tmp_path):
    p = tmp_path / "bad.xlsx"
    write_xlsx(p, HEADERS, [
        ["A", "双液法", "水", 72.8, 21.8, 51.0, 108.0, None, None],
        ["A", "双液法", "二碘甲烷", 50.8, 50.8, 0.0, 200.0, None, None],  # 角度越界
        ["A", "双液法", "", 50.8, 50.8, 0.0, 77.0, None, None],          # 名称缺失
    ])
    preview = parse_measurement_file(p)
    assert len(preview.points) == 1
    assert [e.row_number for e in preview.errors] == [3, 4]


def test_parse_missing_required_column(tmp_path):
    p = tmp_path / "nocol.xlsx"
    write_xlsx(p, ["液体名称", "γ_L(mN/m)"], [["水", 72.8]])
    with pytest.raises(ImportFileError, match="接触角"):
        parse_measurement_file(p)


def test_parse_multiple_methods_error(tmp_path):
    p = tmp_path / "mixed.xlsx"
    write_xlsx(p, HEADERS, [
        ["A", "双液法", "水", 72.8, 21.8, 51.0, 108.0, None, None],
        ["A", "Zisman", "水", 72.8, None, None, 108.0, None, None],
    ])
    preview = parse_measurement_file(p)
    assert preview.method is None
    assert any("多种计算" in e.message for e in preview.errors)


def test_parse_unknown_method(tmp_path):
    p = tmp_path / "badm.xlsx"
    write_xlsx(p, HEADERS, [["A", "神秘方法", "水", 72.8, 21.8, 51.0, 108.0, None, None]])
    preview = parse_measurement_file(p)
    assert any("无法识别的方法" in e.message for e in preview.errors)


def test_parse_method_column_absent(tmp_path):
    headers = ["样品名称", "液体名称", "γ_L(mN/m)", "接触角(°)"]
    p = tmp_path / "nom.xlsx"
    write_xlsx(p, headers, [["A", "水", 72.8, 108.0]])
    preview = parse_measurement_file(p)
    assert preview.method is None
    assert len(preview.points) == 1


def test_excel_report_roundtrip(tmp_path):
    preview = parse_measurement_file(_write_to(tmp_path, "r.xlsx"))
    points = preview.points
    out = tmp_path / "report.xlsx"
    export_excel_report(out, preview.sample_name, METHOD_OWRK, points,
                        preview.remark, preview.test_date)
    wb = openpyxl.load_workbook(out)
    texts = {c.value for row in wb["表面能分析"].iter_rows()
             for c in row if c.value is not None}
    assert "PTFE 样品" in texts
    assert "OWRK 双液法" in texts
    assert any(isinstance(v, str) and v.startswith("总表面能") for v in texts)
    assert any(isinstance(v, str) and "界面能" in v for v in texts)


def test_zisman_report_contains_figure(tmp_path):
    pts = [LiquidPoint(f"液{gamma_l:g}", gamma_l,
                       math.degrees(math.acos(1.4 - 0.02 * gamma_l)))
           for gamma_l in (30.0, 40.0, 50.0)]
    out = tmp_path / "z.xlsx"
    export_excel_report(out, "样品Z", METHOD_ZISMAN, pts)
    wb = openpyxl.load_workbook(out)
    assert len(wb["表面能分析"]._images) == 1
    texts = {c.value for row in wb["表面能分析"].iter_rows()
             for c in row if c.value is not None}
    assert any(isinstance(v, str) and v.startswith("临界表面张力") for v in texts)


def test_pdf_report(tmp_path):
    preview = parse_measurement_file(_write_to(tmp_path, "p.xlsx"))
    out = tmp_path / "report.pdf"
    export_pdf_report(out, preview.sample_name, METHOD_OWRK, preview.points,
                      preview.remark, preview.test_date)
    data = out.read_bytes()
    assert data[:5] == b"%PDF-"
    assert len(data) > 2000


def _write_to(tmp_path, name):
    p = tmp_path / name
    write_xlsx(p, HEADERS, ROWS)
    return p
