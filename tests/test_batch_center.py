import math

import numpy as np
import openpyxl
import pytest

from fiber_tensile.batch_center.core import (
    BatchOptions, discover_data_files, export_results_csv, export_results_excel,
    process_files,
)
from fiber_tensile.data.db import open_database


def _write_xlsx(path, headers, rows):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def test_discover_files_recursive_and_sorted(tmp_path):
    nested = tmp_path / "nested"
    nested.mkdir()
    (nested / "B.csv").write_text("a\n1\n", encoding="utf-8")
    (tmp_path / "a.xlsx").write_bytes(b"x")
    (nested / "ignore.pdf").write_bytes(b"x")

    files = discover_data_files([tmp_path], recursive=True)

    assert [path.name for path in files] == ["a.xlsx", "B.csv"]


def test_fiber_batch_processing_and_duplicate_rename(tmp_path):
    source = tmp_path / "fiber.xlsx"
    _write_xlsx(source, ["批次", "编号", "断裂强力", "强力单位", "直径(μm)"], [
        ["批A", "1", 10, "cN", 5.4],
        ["批A", "2", 11, "cN", 5.4],
    ])
    conn = open_database(str(tmp_path / "test.db"))
    options = BatchOptions("fiber")

    first = process_files(conn, [source], options)[0]
    second = process_files(conn, [source], options)[0]

    assert first.status == "成功" and first.record_name == "批A"
    assert first.valid_rows == 2
    assert second.status == "成功" and second.record_name == "批A (2)"
    assert conn.execute("SELECT COUNT(*) FROM batches").fetchone()[0] == 2
    conn.close()


def test_surface_batch_processing(tmp_path):
    source = tmp_path / "surface.xlsx"
    _write_xlsx(source,
                ["样品名称", "方法", "液体名称", "γ_L(mN/m)", "γ^d(mN/m)",
                 "γ^p(mN/m)", "接触角(°)"], [
                    ["样品A", "OWRK", "水", 72.8, 21.8, 51.0, 80],
                    ["样品A", "OWRK", "二碘甲烷", 50.8, 50.8, 0, 45],
                ])
    conn = open_database(str(tmp_path / "test.db"))

    result = process_files(conn, [source], BatchOptions("surface"))[0]

    assert result.status == "成功"
    assert result.record_name == "样品A"
    assert conn.execute("SELECT COUNT(*) FROM se_measurements").fetchone()[0] == 1
    conn.close()


def test_curve_batch_processing_with_unit_conversion(tmp_path):
    diameter = 5.4
    gauge = 25.0
    area = math.pi * (diameter * 1e-6) ** 2 / 4
    stiffness_n_mm = 230e9 * area / (gauge / 1000) / 1000
    x_mm = np.linspace(0, 0.45, 80)
    force_n = stiffness_n_mm * x_mm
    source = tmp_path / "curve.csv"
    source.write_text(
        "位移,力\n" + "\n".join(
            f"{x * 1000:.8f},{f * 100:.8f}" for x, f in zip(x_mm, force_n)),
        encoding="utf-8-sig")
    conn = open_database(str(tmp_path / "test.db"))
    options = BatchOptions(
        "curve", curve_diameter_um=diameter, curve_gauge_mm=gauge,
        curve_disp_unit="um", curve_force_unit="cN")

    result = process_files(conn, [source], options)[0]

    assert result.status == "成功"
    row = conn.execute("SELECT * FROM fiber_curves").fetchone()
    assert row["e_gpa"] == pytest.approx(230, rel=1e-5)
    conn.close()


def test_failure_isolated_and_exports(tmp_path):
    good = tmp_path / "good.xlsx"
    bad = tmp_path / "bad.xlsx"
    _write_xlsx(good, ["断裂强力", "直径(μm)"], [[10, 5.4]])
    _write_xlsx(bad, ["错误列"], [[1]])
    conn = open_database(str(tmp_path / "test.db"))

    results = process_files(conn, [bad, good], BatchOptions("fiber"))

    assert [result.status for result in results] == ["失败", "成功"]
    csv_path = tmp_path / "log.csv"
    xlsx_path = tmp_path / "log.xlsx"
    export_results_csv(csv_path, results)
    export_results_excel(xlsx_path, results)
    assert csv_path.read_bytes().startswith(b"\xef\xbb\xbf")
    assert openpyxl.load_workbook(xlsx_path)["批量计算结果"].max_row == 3
    conn.close()
