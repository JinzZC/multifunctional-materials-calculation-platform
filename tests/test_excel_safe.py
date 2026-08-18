"""外部文本写入 Excel 时不得被解释为公式。"""

import openpyxl

from fiber_tensile.excel_merger.core import MergeResult, export_merged_workbook
from fiber_tensile.fileio.excel_safe import append_safe_row


def test_append_safe_row_preserves_formula_like_text(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    values = ["=2+2", "+SUM(A1:A2)", "-1+2", "@cmd", "普通文本", 12.5]
    append_safe_row(ws, values)
    path = tmp_path / "safe.xlsx"
    wb.save(path)

    loaded = openpyxl.load_workbook(path, data_only=False)
    cells = list(loaded.active[1])
    assert [cell.value for cell in cells] == values
    assert all(cell.data_type == "s" for cell in cells[:5])
    assert cells[5].data_type == "n"


def test_merged_csv_text_cannot_become_excel_formula(tmp_path):
    result = MergeResult(headers=["样品名", "备注"], rows=[["样品1", "=HYPERLINK(\"x\")"]])
    path = tmp_path / "merged.xlsx"
    export_merged_workbook(path, result)
    loaded = openpyxl.load_workbook(path, data_only=False)
    cell = loaded["合并数据"]["B2"]
    assert cell.value == '=HYPERLINK("x")'
    assert cell.data_type == "s"
