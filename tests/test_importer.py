"""fileio/importer.py 与 fileio/template.py 测试。"""

import openpyxl
import pytest

from fiber_tensile.fileio.importer import parse_decimal_text

from fiber_tensile.core.models import MODE_FILAMENT, MODE_TEXTILE
from fiber_tensile.fileio.importer import (
    ImportFileError, list_sheet_names, parse_file, read_table,
)
from fiber_tensile.fileio.template import generate_template

HEADERS_F = ["编号", "批次", "断裂强力", "强力单位", "直径(μm)", "备注"]
ROWS_F = [
    ["F1", "批A", 10, "cN", 7, "ok"],
    ["F2", "批A", 9.5, "cN", 7.2, None],
    ["F3", "批A", 10.8, "cN", 6.8, None],
]

HEADERS_T = ["编号", "批次", "断裂强力", "线密度(dtex)", "密度(g/cm³)", "测试日期"]
ROWS_T = [
    ["T1", "批B", 12.5, 1.5, 1.38, "2026-08-14"],
    ["T2", "批B", 11.8, 1.55, 1.38, "2026/8/15"],
]


def write_xlsx(path, headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    wb.save(path)


def write_csv(path, headers, rows, encoding="utf-8-sig", sep=","):
    with open(path, "w", encoding=encoding, newline="") as f:
        f.write(sep.join(headers) + "\n")
        for r in rows:
            f.write(sep.join("" if v is None else str(v) for v in r) + "\n")


# ---------- 基本解析：三种格式 ----------

def test_xlsx_filament_basic(tmp_path):
    p = tmp_path / "data.xlsx"
    write_xlsx(p, HEADERS_F, ROWS_F)
    preview = parse_file(p)
    assert preview.inferred_mode == MODE_FILAMENT
    assert preview.valid_rows == 3
    assert preview.errors == []
    assert preview.batch_from_file == "批A"
    assert preview.samples[0].diameter_um == 7.0
    assert preview.samples[0].breaking_force == 10.0
    assert preview.samples[0].force_unit == "cN"
    assert preview.samples[0].linear_density_dtex is None


def test_csv_utf8_sig(tmp_path):
    p = tmp_path / "data.csv"
    write_csv(p, HEADERS_F, ROWS_F, encoding="utf-8-sig")
    preview = parse_file(p)
    assert preview.valid_rows == 3


def test_csv_gbk(tmp_path):
    p = tmp_path / "data_gbk.csv"
    write_csv(p, HEADERS_F, ROWS_F, encoding="gb18030")
    preview = parse_file(p)
    assert preview.valid_rows == 3


def test_txt_tab_delimited(tmp_path):
    p = tmp_path / "data.txt"
    write_csv(p, HEADERS_F, ROWS_F, sep="\t")
    preview = parse_file(p)
    assert preview.valid_rows == 3


def test_tsv_tab_delimited(tmp_path):
    p = tmp_path / "data.tsv"
    write_csv(p, HEADERS_F, ROWS_F, sep="\t")
    preview = parse_file(p)
    assert preview.valid_rows == 3


def test_auto_header_skips_preamble_and_preserves_source_row(tmp_path):
    p = tmp_path / "instrument.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["仪器导出报告"])
    ws.append(["操作者", "张三"])
    ws.append(HEADERS_F)
    ws.append(["F1", "批A", "bad", "cN", 7, None])
    wb.save(p)
    preview = parse_file(p)
    assert preview.header_row == 3
    assert preview.valid_rows == 0
    assert preview.errors[0].row_number == 4


def test_select_excel_sheet_and_manual_header(tmp_path):
    p = tmp_path / "sheets.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "说明"
    wb.active.append(["请勿导入"])
    ws = wb.create_sheet("测试数据")
    ws.append(["标题"])
    ws.append(HEADERS_F)
    ws.append(ROWS_F[0])
    wb.save(p)
    assert list_sheet_names(p) == ["说明", "测试数据"]
    preview = parse_file(p, sheet_name="测试数据", header_row=2)
    assert preview.valid_rows == 1
    assert preview.header_row == 2
    assert preview.sheet_name == "测试数据"


def test_explicit_column_mapping_and_header_force_unit(tmp_path):
    p = tmp_path / "custom.xlsx"
    write_xlsx(p, ["Specimen", "Peak Load(N)", "Dia micron"],
               [["A-1", 0.12, 7.0]])
    preview = parse_file(
        p, mode=MODE_FILAMENT,
        column_mapping={
            "sample_no": "Specimen",
            "breaking_force": "Peak Load(N)",
            "diameter_um": "Dia micron",
        },
    )
    assert preview.valid_rows == 1
    assert preview.samples[0].force_unit == "N"
    assert preview.samples[0].breaking_force == 0.12


def test_duplicate_sample_number_warns(tmp_path):
    p = tmp_path / "duplicate.xlsx"
    rows = [ROWS_F[0], [" f1 ", "批A", 9.0, "cN", 7.0, None]]
    write_xlsx(p, HEADERS_F, rows)
    preview = parse_file(p)
    assert preview.valid_rows == 2
    assert any("重复试样" in warning.message for warning in preview.warnings)


def test_textile_mode_with_dates(tmp_path):
    p = tmp_path / "t.xlsx"
    write_xlsx(p, HEADERS_T, ROWS_T)
    preview = parse_file(p)
    assert preview.inferred_mode == MODE_TEXTILE
    assert preview.valid_rows == 2
    assert preview.samples[0].linear_density_dtex == 1.5
    assert preview.samples[0].test_date == "2026-08-14"
    assert preview.samples[1].test_date == "2026-08-15"


# ---------- 模式指定与冲突 ----------

def test_explicit_mode_overrides(tmp_path):
    p = tmp_path / "t.xlsx"
    write_xlsx(p, HEADERS_T, ROWS_T)
    preview = parse_file(p, mode=MODE_TEXTILE)
    assert preview.valid_rows == 2
    assert preview.samples[0].diameter_um is None


def test_mode_mismatch_raises(tmp_path):
    p = tmp_path / "t.xlsx"
    write_xlsx(p, HEADERS_T, ROWS_T)
    with pytest.raises(ImportFileError, match="直径"):
        parse_file(p, mode=MODE_FILAMENT)


def test_both_mode_columns_raises(tmp_path):
    p = tmp_path / "both.xlsx"
    write_xlsx(p, ["断裂强力", "直径(μm)", "线密度(dtex)"], [[10, 7, 1.5]])
    with pytest.raises(ImportFileError, match="同时"):
        parse_file(p)


def test_no_mode_column_raises(tmp_path):
    p = tmp_path / "none.xlsx"
    write_xlsx(p, ["断裂强力"], [[10]])
    with pytest.raises(ImportFileError, match="直径"):
        parse_file(p)


def test_missing_force_column_raises(tmp_path):
    p = tmp_path / "noforce.xlsx"
    write_xlsx(p, ["编号", "直径(μm)"], [["F1", 7]])
    with pytest.raises(ImportFileError, match="断裂强力"):
        parse_file(p)


# ---------- 表头别名与容错 ----------

def test_header_aliases(tmp_path):
    # 全角括号、大小写、别名混合
    p = tmp_path / "alias.xlsx"
    write_xlsx(p, ["Sample_No", "批次号", "Force", "UNIT", "直径（μm）"], [["X1", "批X", 10, "N", 7]])
    preview = parse_file(p)
    assert preview.valid_rows == 1
    s = preview.samples[0]
    assert s.sample_no == "X1"
    assert s.force_unit == "N"
    assert s.diameter_um == 7.0


# ---------- 数值解析 ----------

def test_full_width_numbers_and_thousands_sep(tmp_path):
    p = tmp_path / "num.xlsx"
    write_xlsx(p, HEADERS_F, [["F1", None, "１,２３４.５", None, 7]])
    preview = parse_file(p)
    assert preview.samples[0].breaking_force == 1234.5


def test_force_unit_defaults_to_cn(tmp_path):
    p = tmp_path / "unit.xlsx"
    write_xlsx(p, ["断裂强力", "直径(μm)"], [[10, 7]])
    preview = parse_file(p)
    assert preview.samples[0].force_unit == "cN"


# ---------- 错误行定位 ----------

def test_error_rows_reported_with_line_numbers(tmp_path):
    p = tmp_path / "err.xlsx"
    write_xlsx(p, HEADERS_F, [
        ["F1", "批A", 10, "cN", 7, None],        # 行 2：有效
        ["F2", "批A", "abc", "cN", 7, None],     # 行 3：强力无效
        ["F3", "批A", 10, "cN", 0, None],        # 行 4：直径非正
        ["F4", "批A", 10, "kg", 7, None],        # 行 5：单位无效
    ])
    preview = parse_file(p)
    assert preview.valid_rows == 1
    assert [e.row_number for e in preview.errors] == [3, 4, 5]
    assert any("不是有效数字" in e.message for e in preview.errors if e.row_number == 3)
    assert any("必须为正数" in e.message for e in preview.errors if e.row_number == 4)
    assert any("单位" in e.message for e in preview.errors if e.row_number == 5)


def test_all_empty_rows_skipped(tmp_path):
    p = tmp_path / "empty.xlsx"
    write_xlsx(p, HEADERS_F, [["F1", "批A", 10, "cN", 7, None], [None, None, None, None, None, None]])
    preview = parse_file(p)
    assert preview.valid_rows == 1
    assert preview.errors == []


# ---------- 日期解析 ----------

def test_date_excel_serial(tmp_path):
    # Excel 序列号 45601 = 2024-11-04 附近；直接用已知锚点：45658 = 2025-01-01
    p = tmp_path / "d.xlsx"
    write_xlsx(p, HEADERS_T, [["T1", "批B", 10, 1.5, None, 45658]])
    preview = parse_file(p)
    assert preview.samples[0].test_date == "2025-01-01"


def test_bad_date_warns_and_clears(tmp_path):
    p = tmp_path / "d.xlsx"
    write_xlsx(p, HEADERS_T, [["T1", "批B", 10, 1.5, None, "明天"]])
    preview = parse_file(p)
    assert preview.valid_rows == 1          # 行仍有效
    assert preview.samples[0].test_date is None
    assert any("日期" in w.message for w in preview.warnings)


# ---------- 范围警告 ----------

def test_out_of_range_warns_not_blocks(tmp_path):
    p = tmp_path / "r.xlsx"
    write_xlsx(p, HEADERS_F, [["F1", "批A", 10, "cN", 2000, None]])  # 直径超 1000 μm
    preview = parse_file(p)
    assert preview.valid_rows == 1
    assert any("直径" in w.message and "范围" in w.message for w in preview.warnings)


# ---------- 文件级错误 ----------

def test_unsupported_extension_raises(tmp_path):
    p = tmp_path / "data.docx"
    p.write_bytes(b"x")
    with pytest.raises(ImportFileError, match="格式"):
        parse_file(p)


def test_xls_friendly_error(tmp_path):
    p = tmp_path / "data.xls"
    p.write_bytes(b"x")
    with pytest.raises(ImportFileError, match="另存为"):
        parse_file(p)


def test_missing_file_raises(tmp_path):
    with pytest.raises(ImportFileError, match="不存在"):
        parse_file(tmp_path / "nope.xlsx")


# ---------- 模板 ----------

def test_parse_decimal_text():
    # 欧式小数逗号（逗号后 1-2 位数字）
    assert parse_decimal_text("0,5") == "0.5"
    assert parse_decimal_text("1,5") == "1.5"
    assert parse_decimal_text("12,34") == "12.34"
    # 千分位逗号（逗号后三位数字）
    assert parse_decimal_text("1,234") == "1234"
    assert parse_decimal_text("１２,３４５") == "12345"    # 全角归一
    # 普通数字不受影响
    assert parse_decimal_text("72.8") == "72.8"


def test_template_generated_and_reimportable(tmp_path):
    p = tmp_path / "模板.xlsx"
    generate_template(p)

    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["单丝模式示例", "纺织模式示例", "填写说明"]
    # 模板第一张表是单丝示例，应可直接解析
    preview = parse_file(p)
    assert preview.inferred_mode == MODE_FILAMENT
    assert preview.valid_rows == 3
    assert preview.errors == []
