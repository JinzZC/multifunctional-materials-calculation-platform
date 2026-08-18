"""Excel 多文件夹提取与合并模块测试。"""

from pathlib import Path

import openpyxl
import pytest

from fiber_tensile.excel_merger.core import (
    MergeCancelled, MergeOptions, MergeResult, discover_excel_files, export_merged_workbook,
    merge_workbooks,
)


def _write(path, sheet, headers, rows, title_row=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    if title_row is not None:
        ws.append([title_row])
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(path)


def test_discover_multiple_folders_recursive_and_deduplicate(tmp_path):
    a = tmp_path / "甲"
    b = tmp_path / "乙" / "子目录"
    a.mkdir()
    b.mkdir(parents=True)
    _write(a / "a.xlsx", "数据", ["编号"], [[1]])
    _write(b / "b.xlsm", "数据", ["编号"], [[2]])
    (b / "c.csv").write_text("编号,值\n3,30\n", encoding="utf-8-sig")
    (b / "d.txt").write_text("编号\t值\n4\t40\n", encoding="utf-8")
    _write(b / "~$锁文件.xlsx", "数据", ["编号"], [[3]])
    assert sorted(p.name for p in discover_excel_files([a, a, tmp_path / "乙"])) == [
        "a.xlsx", "b.xlsm", "c.csv", "d.txt"]
    assert [p.name for p in discover_excel_files([tmp_path / "乙"], recursive=False)] == []


def test_merge_union_aligns_normalised_headers_and_tracks_source(tmp_path):
    one = tmp_path / "一.xlsx"
    two = tmp_path / "二.xlsx"
    _write(one, "数据", ["样品编号", " 强度 (GPa) "], [["A", 3.1], ["B", 3.2]])
    _write(two, "结果", ["样品编号", "强度(GPa)", "备注"], [["C", 3.3, "复测"]])
    result = merge_workbooks([one, two], MergeOptions(), [tmp_path])
    assert result.headers == ["来源文件夹", "来源文件", "来源工作表", "原始行号",
                              "样品编号", "强度 (GPa)", "备注"]
    assert len(result.rows) == 3
    assert result.rows[0][-3:] == ["A", 3.1, None]
    assert result.rows[2][-3:] == ["C", 3.3, "复测"]
    assert result.rows[2][1] == "二.xlsx"
    assert result.rows[2][2] == "结果"
    assert result.rows[2][3] == 2


def test_header_row_specific_sheet_intersection_and_empty_rows(tmp_path):
    one = tmp_path / "一.xlsx"
    two = tmp_path / "二.xlsx"
    _write(one, "测试数据", ["编号", "强度", "备注"], [[1, 10, "a"], [None, None, None]],
           title_row="仪器报告")
    _write(two, "测试数据", ["编号", "强度", "温度"], [[2, 11, 25]], title_row="报告")
    result = merge_workbooks(
        [one, two], MergeOptions(header_row=2, sheet_name="测试数据",
                                 include_source=False, column_mode="intersection"))
    assert result.headers == ["编号", "强度"]
    assert result.rows == [[1, 10], [2, 11]]


def test_missing_sheet_and_corrupt_file_are_logged(tmp_path):
    good = tmp_path / "good.xlsx"
    bad = tmp_path / "bad.xlsx"
    _write(good, "其他表", ["编号"], [[1]])
    bad.write_bytes(b"not an Excel file")
    result = merge_workbooks([good, bad], MergeOptions(sheet_name="目标表"))
    assert result.rows == []
    assert {s.status for s in result.statuses} == {"跳过", "失败"}
    assert any("工作表不存在" in s.message for s in result.statuses)
    assert any("无法打开" in s.message for s in result.statuses)


def test_duplicate_and_blank_headers_are_renamed(tmp_path):
    path = tmp_path / "dup.xlsx"
    _write(path, "数据", ["编号", "编号", None], [[1, 2, 3]])
    result = merge_workbooks([path], MergeOptions(include_source=False))
    assert result.headers == ["编号", "编号_2"]
    # 第三个空表头位于有效表头范围之外，不会凭空扩展数据范围。
    assert len(result.warnings) == 1
    assert result.rows == [[1, 2]]


def test_export_contains_data_and_audit_log(tmp_path):
    source = tmp_path / "source.xlsx"
    _write(source, "数据", ["编号", "值"], [[1, 2]])
    result = merge_workbooks([source])
    output = tmp_path / "总表.xlsx"
    export_merged_workbook(output, result)
    wb = openpyxl.load_workbook(output, data_only=True)
    assert wb.sheetnames == ["合并数据", "合并日志"]
    assert wb["合并数据"]["E2"].value == 1
    assert wb["合并数据"].freeze_panes == "A2"
    assert wb["合并日志"]["A2"].value == "成功"


def test_scan_excludes_explicit_output_path(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "总表.xlsx"
    _write(source, "数据", ["编号"], [[1]])
    _write(output, "数据", ["编号"], [[2]])
    found = discover_excel_files([tmp_path], excluded_paths=[output])
    assert found == [source.resolve()]


def test_merge_csv_txt_encodings_and_delimiters(tmp_path):
    csv_path = tmp_path / "数据.csv"
    txt_path = tmp_path / "数据.txt"
    csv_path.write_bytes("样品名称,强度,备注\n样品10,3.0,常规\n".encode("gb18030"))
    txt_path.write_bytes("样品名称\t强度\n样品2\t3.2\n".encode("utf-16"))
    result = merge_workbooks(
        [csv_path, txt_path], MergeOptions(include_source=False, sort_by_sample=True))
    assert result.headers == ["样品名称", "强度", "备注"]
    assert result.rows == [["样品2", "3.2", None], ["样品10", "3.0", "常规"]]
    assert all(status.status == "成功" for status in result.statuses)
    assert "gb18030" in result.statuses[0].message
    assert "utf-16" in result.statuses[1].message


def test_semicolon_csv_and_explicit_sample_column(tmp_path):
    path = tmp_path / "semicolon.csv"
    path.write_text("材料编号;测量值\nB-12;1\nB-3;2\n", encoding="utf-8-sig")
    result = merge_workbooks(
        [path], MergeOptions(include_source=False, sort_by_sample=True,
                             sample_column="材料编号"))
    assert result.rows == [["B-3", "2"], ["B-12", "1"]]
    assert "分号分隔" in result.statuses[0].message


def test_sort_requires_recognisable_sample_column(tmp_path):
    path = tmp_path / "unknown.csv"
    path.write_text("编号,值\nA,1\n", encoding="utf-8")
    with pytest.raises(ValueError, match="样品名列"):
        merge_workbooks([path], MergeOptions(sort_by_sample=True))


def test_tsv_and_quoted_multiline_csv_are_not_split(tmp_path):
    csv_path = tmp_path / "quoted.csv"
    tsv_path = tmp_path / "more.tsv"
    csv_path.write_text('样品名,备注\n样品2,"第一行\n第二行"\n', encoding="utf-8-sig")
    tsv_path.write_text("样品名\t备注\n样品1\t正常\n", encoding="utf-8-sig")
    result = merge_workbooks(
        [csv_path, tsv_path], MergeOptions(include_source=False, sort_by_sample=True))
    assert result.rows[0] == ["样品1", "正常"]
    assert result.rows[1][0] == "样品2"
    assert result.rows[1][1].splitlines() == ["第一行", "第二行"]


def test_sheet_name_contains_extracts_all_matches(tmp_path):
    path = tmp_path / "multi.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "样品数据_A"
    wb.active.append(["样品名", "值"])
    wb.active.append(["样品2", 2])
    ws = wb.create_sheet("历史样品数据_B")
    ws.append(["样品名", "值"])
    ws.append(["样品1", 1])
    ws = wb.create_sheet("汇总")
    ws.append(["样品名", "值"])
    ws.append(["不应提取", 99])
    wb.save(path)

    result = merge_workbooks(
        [path], MergeOptions(sheet_name="样品数据", sheet_match_mode="contains",
                             include_source=False, sort_by_sample=True))
    assert result.rows == [["样品1", 1], ["样品2", 2]]
    assert [status.sheet for status in result.statuses] == ["样品数据_A", "历史样品数据_B"]


def test_sheet_contains_normalisation_no_match_and_csv_unaffected(tmp_path):
    book = tmp_path / "book.xlsx"
    csv_path = tmp_path / "data.csv"
    _write(book, "ＴＥＳＴ结果", ["样品名", "值"], [["样品2", 2]])
    csv_path.write_text("样品名,值\n样品1,1\n", encoding="utf-8-sig")
    result = merge_workbooks(
        [book, csv_path], MergeOptions(sheet_name="test", sheet_match_mode="contains",
                                       include_source=False, sort_by_sample=True))
    assert [row[0] for row in result.rows] == ["样品1", "样品2"]

    no_match = merge_workbooks(
        [book], MergeOptions(sheet_name="不存在", sheet_match_mode="contains"))
    assert no_match.rows == []
    assert no_match.statuses[0].status == "跳过"
    assert "没有名称包含" in no_match.statuses[0].message


def test_filename_contains_filter_is_recursive_and_format_independent(tmp_path):
    nested = tmp_path / "子目录"
    nested.mkdir()
    _write(tmp_path / "拉伸_甲.xlsx", "数据", ["样品名"], [["A"]])
    (nested / "拉伸_乙.csv").write_text("样品名\nB\n", encoding="utf-8")
    (nested / "接触角.txt").write_text("样品名\nC\n", encoding="utf-8")
    found = discover_excel_files([tmp_path], filename_contains="拉伸")
    assert sorted(path.name for path in found) == ["拉伸_乙.csv", "拉伸_甲.xlsx"]


def test_filename_filter_ignores_case_and_fullwidth(tmp_path):
    _write(tmp_path / "ＴＥＳＴ_Result.xlsx", "数据", ["样品名"], [["A"]])
    _write(tmp_path / "other.xlsx", "数据", ["样品名"], [["B"]])
    found = discover_excel_files([tmp_path], filename_contains="test")
    assert [path.name for path in found] == ["ＴＥＳＴ_Result.xlsx"]


def test_formula_without_cached_value_is_preserved_and_warned(tmp_path):
    path = tmp_path / "formula.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["样品名", "计算值"])
    ws.append(["样品1", "=1+2"])
    wb.save(path)

    result = merge_workbooks(
        [path], MergeOptions(include_source=False, formula_mode="values"))
    assert result.rows == [["样品1", "=1+2"]]
    assert any("没有缓存结果" in warning and "1 个公式单元格" in warning
               for warning in result.warnings)


def test_formula_text_mode_preserves_formula_without_cache_warning(tmp_path):
    path = tmp_path / "formula.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["样品名", "计算值"])
    ws.append(["样品1", "=SUM(1,2)"])
    wb.save(path)

    result = merge_workbooks(
        [path], MergeOptions(include_source=False, formula_mode="formulas"))
    assert result.rows == [["样品1", "=SUM(1,2)"]]
    assert not any("没有缓存结果" in warning for warning in result.warnings)


def test_column_mapping_unifies_instrument_headers(tmp_path):
    one = tmp_path / "one.csv"
    two = tmp_path / "two.csv"
    one.write_text("试样名称,值\n样品2,2\n", encoding="utf-8")
    two.write_text("Sample Name,值\n样品1,1\n", encoding="utf-8")
    result = merge_workbooks([one, two], MergeOptions(
        include_source=False, sort_by_sample=True,
        column_mappings=(("试样名称", "样品名"), ("Sample Name", "样品名"))))
    assert result.headers == ["样品名", "值"]
    assert result.rows == [["样品1", "1"], ["样品2", "2"]]
    assert sum("字段" in warning and "已统一" in warning for warning in result.warnings) == 2


def test_duplicate_detection_keep_first_and_last(tmp_path):
    path = tmp_path / "dup.csv"
    path.write_text("样品名,日期,值\nA,2026-01-01,旧\nA,2026-01-01,新\nB,2026-01-01,唯一\n",
                    encoding="utf-8")
    common = dict(include_source=False, sort_by_sample=True,
                  duplicate_columns=("样品名", "日期"))
    first = merge_workbooks([path], MergeOptions(**common, duplicate_mode="first"))
    last = merge_workbooks([path], MergeOptions(**common, duplicate_mode="last"))
    assert first.rows == [["A", "2026-01-01", "旧"], ["B", "2026-01-01", "唯一"]]
    assert last.rows == [["A", "2026-01-01", "新"], ["B", "2026-01-01", "唯一"]]
    assert any("移除 1 行" in warning for warning in first.warnings)


def test_keep_all_reports_duplicates_and_multi_column_sort(tmp_path):
    path = tmp_path / "sort.csv"
    path.write_text("样品名,测试序号,值\n样品2,10,A\n样品1,2,B\n样品1,1,C\n样品1,1,D\n",
                    encoding="utf-8")
    result = merge_workbooks([path], MergeOptions(
        include_source=False, sort_by_sample=True, sort_columns=("测试序号",),
        duplicate_mode="all", duplicate_columns=("样品名", "测试序号")))
    assert [(row[0], row[1]) for row in result.rows] == [
        ("样品1", "1"), ("样品1", "1"), ("样品1", "2"), ("样品2", "10")]
    assert any("发现 1 条重复记录，已全部保留" in warning for warning in result.warnings)


def test_progress_callback_and_cancellation(tmp_path):
    paths = []
    for index in range(3):
        path = tmp_path / f"{index}.csv"
        path.write_text("样品名,值\nA,1\n", encoding="utf-8")
        paths.append(path)
    progress = []
    result = merge_workbooks(paths, progress_callback=lambda *args: progress.append(args))
    assert len(result.rows) == 3
    assert progress[0][:2] == (0, 3)
    assert progress[-1][:2] == (3, 3)

    checks = iter([False, True])
    with pytest.raises(MergeCancelled):
        merge_workbooks(paths, is_cancelled=lambda: next(checks, True))


def test_export_splits_rows_across_multiple_data_sheets(tmp_path):
    result = MergeResult(headers=["样品名", "值"],
                         rows=[[f"样品{i}", i] for i in range(5)])
    output = tmp_path / "split.xlsx"
    export_merged_workbook(output, result, max_rows_per_sheet=3)  # 每表 1 行表头 + 2 行数据
    wb = openpyxl.load_workbook(output, data_only=True)
    assert wb.sheetnames == ["合并数据", "合并数据_2", "合并数据_3", "合并日志"]
    assert [wb[name].max_row for name in wb.sheetnames[:3]] == [3, 3, 2]
    assert "拆分为 3 个" in wb["合并日志"]["F2"].value
