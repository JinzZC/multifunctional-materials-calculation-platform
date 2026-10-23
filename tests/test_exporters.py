"""fileio/exporter_excel.py 与 fileio/exporter_pdf.py 测试。"""

import math

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pytest

from fiber_tensile.core.models import MODE_FILAMENT, MODE_TEXTILE, Batch, Sample
from fiber_tensile.core.statistics import batch_statistics
from fiber_tensile.core.weibull import weibull_fit, weibull_fit_mle, weibull_fit_3p, weibull_in_gpa
from fiber_tensile.data.db import open_database
from fiber_tensile.data.repository import add_samples, create_batch, get_batch
from fiber_tensile.fileio.exporter_excel import export_excel_report
from fiber_tensile.fileio.exporter_pdf import export_pdf_report


@pytest.fixture
def db_conn(tmp_path):
    conn = open_database(str(tmp_path / "t.db"))
    yield conn
    conn.close()


@pytest.fixture
def filament_batch(db_conn):
    bid = create_batch(db_conn, Batch(batch_name="碳纤维-A", mode=MODE_FILAMENT,
                                      fiber_type="T300"))
    add_samples(db_conn, bid, [
        Sample(breaking_force=f, diameter_um=d, sample_no=f"F{i}")
        for i, (f, d) in enumerate(
            [(10.0, 7.0), (9.5, 7.2), (10.8, 6.8), (9.0, 7.1), (10.2, 6.9)], start=1)])
    return get_batch(db_conn, bid)


def make_analysis(batch):
    strengths = [s.strength for s in batch.samples]
    return batch_statistics(strengths), weibull_fit(strengths)


def _cell_texts(ws):
    return {c.value for row in ws.iter_rows() for c in row if c.value is not None}


def test_excel_report_roundtrip(tmp_path, filament_batch):
    analysis, weibull = make_analysis(filament_batch)
    out = tmp_path / "report.xlsx"
    export_excel_report(out, filament_batch, analysis, weibull)

    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["统计汇总", "数据明细", "Weibull 分析数据", "Weibull 拟合直线数据"]

    texts = _cell_texts(wb["统计汇总"])
    assert "碳纤维-A" in texts
    assert "单丝模式" in texts
    assert "T300" in texts
    assert any(isinstance(v, float) and v == pytest.approx(weibull.m, rel=1e-4) for v in texts)

    ws2 = wb["数据明细"]
    assert ws2["D2"].value == pytest.approx(7.0)                      # 直径列
    assert ws2["F2"].value == pytest.approx(2598.5, rel=1e-3)         # 强度列
    assert len(wb["统计汇总"]._images) == 1                            # Weibull 图已插入

    ws3 = wb["Weibull 分析数据"]
    assert ws3["C2"].value == pytest.approx(0.1)                      # n=5 时第一个 median rank


def test_excel_report_records_origin_rank_method(tmp_path, filament_batch):
    strengths = [s.strength for s in filament_batch.samples]
    analysis = batch_statistics(strengths)
    weibull = weibull_fit(strengths, "mean_rank")
    out = tmp_path / "origin-rank.xlsx"
    export_excel_report(out, filament_batch, analysis, weibull)

    wb = openpyxl.load_workbook(out)
    summary = _cell_texts(wb["统计汇总"])
    assert "Origin 兼容（Mean Rank / Van der Waerden）" in summary
    assert "Origin 兼容" in wb["Weibull 分析数据"]["C1"].value
    assert wb["Weibull 分析数据"]["C2"].value == pytest.approx(1 / 6, abs=5e-7)
    assert wb["Weibull 分析数据"]["E1"].value == "ln(−ln(1−F̂))"


def test_excel_report_single_sample_no_weibull(tmp_path, db_conn):
    bid = create_batch(db_conn, Batch(batch_name="单根", mode=MODE_FILAMENT))
    add_samples(db_conn, bid, [Sample(breaking_force=10.0, diameter_um=7.0)])
    b = get_batch(db_conn, bid)
    analysis = batch_statistics([s.strength for s in b.samples])

    out = tmp_path / "single.xlsx"
    export_excel_report(out, b, analysis, None)

    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["统计汇总", "数据明细"]
    texts = _cell_texts(wb["统计汇总"])
    assert any("至少需要 2 个" in str(v) for v in texts)
    assert len(wb["统计汇总"]._images) == 0


def test_textile_report_with_density_units(tmp_path, db_conn):
    bid = create_batch(db_conn, Batch(batch_name="涤纶-B", mode=MODE_TEXTILE, density=1.38))
    add_samples(db_conn, bid, [
        Sample(breaking_force=f, linear_density_dtex=t, sample_no=f"T{i}")
        for i, (f, t) in enumerate([(12.5, 1.5), (11.8, 1.55), (13.1, 1.48)], start=1)])
    b = get_batch(db_conn, bid)
    analysis, weibull = make_analysis(b)

    out = tmp_path / "t.xlsx"
    export_excel_report(out, b, analysis, weibull)

    wb = openpyxl.load_workbook(out)
    texts = _cell_texts(wb["统计汇总"])
    assert "纺织纤维模式" in texts
    assert any("cN/dtex" in str(v) for v in texts)
    assert "特征强度 σ₀(GPa)" in texts
    data = list(wb["Weibull 分析数据"].values)
    expected_strengths = sorted(s.strength * 1.38 / 10 for s in b.samples)
    assert [row[1] for row in data[1:]] == pytest.approx(expected_strengths)
    assert [row[3] for row in data[1:]] == pytest.approx(np.log(expected_strengths))


def test_pdf_report(tmp_path, filament_batch):
    analysis, weibull = make_analysis(filament_batch)
    out = tmp_path / "report.pdf"
    export_pdf_report(out, filament_batch, analysis, weibull)

    data = out.read_bytes()
    assert data[:5] == b"%PDF-"
    assert len(data) > 10_000      # 含中文 CID 字体子集与嵌入图片


def test_pdf_report_single_sample(tmp_path, db_conn):
    bid = create_batch(db_conn, Batch(batch_name="单根", mode=MODE_TEXTILE))
    add_samples(db_conn, bid, [Sample(breaking_force=10.0, linear_density_dtex=2.0)])
    b = get_batch(db_conn, bid)
    analysis = batch_statistics([s.strength for s in b.samples])

    out = tmp_path / "single.pdf"
    export_pdf_report(out, b, analysis, None)
    assert out.read_bytes()[:5] == b"%PDF-"


@pytest.mark.parametrize("fit", [weibull_fit, weibull_fit_mle, weibull_fit_3p])
def test_exported_gpa_scatter_and_fit_line_match_figure(tmp_path, filament_batch, fit):
    from fiber_tensile.plotting import build_weibull_figure

    strengths = [s.strength for s in filament_batch.samples]
    raw = fit(strengths, "mean_rank")
    out = tmp_path / "plot-data.xlsx"
    export_excel_report(out, filament_batch, batch_statistics(strengths), raw)
    wb = openpyxl.load_workbook(out)
    scatter = list(wb["Weibull 分析数据"].values)
    assert scatter[0][1] == "σ(GPa)"
    assert scatter[0][3] == ("ln((σ−γ)/GPa)" if raw.gamma > 0 else "ln(σ/GPa)")
    # 完整精度写入，避免原来的两位强度/六位对数截断。
    assert [row[1] for row in scatter[1:]] == pytest.approx([s / 1000 for s in sorted(strengths)], rel=1e-14)
    gpa = weibull_in_gpa(raw)
    fig = build_weibull_figure([raw], sigma0_unit="MPa")
    try:
        ax = fig.axes[0]
        assert np.asarray([[row[3], row[4]] for row in scatter[1:]]) == pytest.approx(
            np.asarray(ax.collections[0].get_offsets()), rel=1e-14)
        line = np.asarray(list(wb["Weibull 拟合直线数据"].iter_rows(
            min_row=2, max_row=101, min_col=1, max_col=2, values_only=True)))
        assert line.shape == (100, 2)
        assert line[:, 0] == pytest.approx(ax.lines[0].get_xdata(), rel=1e-14)
        assert line[:, 1] == pytest.approx(ax.lines[0].get_ydata(), rel=1e-14)
        assert [row[5] for row in scatter[1:]] == pytest.approx(
            gpa.m * np.asarray(gpa.x) - gpa.m * math.log(gpa.sigma0), rel=1e-14)
        parameters = dict(wb["Weibull 拟合直线数据"].iter_rows(
            min_row=1, max_row=11, min_col=4, max_col=5, values_only=True))
        assert parameters["斜率 m"] == pytest.approx(raw.m, rel=1e-14)
        assert parameters["截距 a"] == pytest.approx(-raw.m * math.log(raw.sigma0 / 1000), rel=1e-14)
    finally:
        plt.close(fig)
        wb.close()


def test_export_already_gpa_result_is_not_converted_twice(tmp_path, filament_batch):
    stats, raw = make_analysis(filament_batch)
    out = tmp_path / "already-gpa.xlsx"
    export_excel_report(out, filament_batch, stats, weibull_in_gpa(raw))
    wb = openpyxl.load_workbook(out)
    assert wb["Weibull 分析数据"]["B2"].value == pytest.approx(raw.strengths_sorted[0] / 1000)
    wb.close()


def test_textile_without_density_does_not_export_false_gpa(tmp_path, db_conn):
    bid = create_batch(db_conn, Batch(batch_name="缺少密度", mode=MODE_TEXTILE))
    add_samples(db_conn, bid, [Sample(breaking_force=f, linear_density_dtex=2) for f in (10, 12, 14)])
    batch = get_batch(db_conn, bid)
    stats, raw = make_analysis(batch)
    out = tmp_path / "no-density.xlsx"
    export_excel_report(out, batch, stats, raw)
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["统计汇总", "数据明细"]
    assert any("需填写有效密度" in str(v) for v in _cell_texts(wb["统计汇总"]))
    assert wb["数据明细"]["F1"].value == "强度(cN/dtex)"
    wb.close()
